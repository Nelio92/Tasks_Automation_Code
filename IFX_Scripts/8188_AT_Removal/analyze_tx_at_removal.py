from __future__ import annotations
import argparse
import csv
import math
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from statistics import NormalDist
from time import perf_counter
from typing import Any
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


INPUT_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = INPUT_DIR / "Output"
REPORT_PATH = OUTPUT_DIR / "AT_Removal_Module_Analysis.xlsx"

DEFAULT_MODULES = ("DPLL", "TXGE", "TXLO", "TXPA", "TXPB", "TXPC", "TXPD", "TXPS")
DEFAULT_MONITORING_PATTERNS = ("LowPowDtct", "OpDtct", "LckDtct", "OpErr", "_S980", "SuFuId", "DtsTx", "DtsRx")

# Direct-run inputs used when launching the script without CLI arguments.
RUN_MODULES = DEFAULT_MODULES
RUN_MAX_TESTS_PER_MODULE = None
RUN_SKIP_PLOTS = False
RUN_MONITORING_PATTERNS = DEFAULT_MONITORING_PATTERNS

META_COLS = (
    "WAFER",
    "X",
    "Y",
    "SITE_NUM",
    "PF",
    "HBIN",
    "SBIN",
    "CHIP_ID",
    "FIRST_FAIL_TEST",
)
HEADER_ROWS = 13
DATA_START_ROW = 18
ENCODING = "latin1"
DELIMITER = ";"
PROBABILITY_TICKS_PCT = (0.01, 0.1, 1.0, 10.0, 50.0, 90.0, 99.0, 99.9, 99.99)
INSERTION_ORDER = ("S1", "S2", "S3")
MAX_CDF_PLOT_POINTS = 1200
HIGH_CORRELATION_R2_THRESHOLD = 0.8
LOW_ESCAPE_PPM_THRESHOLD = 1.0
SIGMA_SEARCH_VALUES = tuple(step / 2.0 for step in range(0, 25))
ZERO_ESCAPE_CONFIDENCE_LEVELS = (0.90, 0.95, 0.99)
INSERTION_LABELS = {
    "S1": "S1 hot",
    "S2": "S2 cold",
    "S3": "S3 ambient",
}
INSERTION_COLORS = {
    "S1": "#C00000",
    "S2": "#0070C0",
    "S3": "#00A65A",
}
INSERTION_LIMIT_STYLES = {
    "S1": "-",
    "S2": "--",
    "S3": ":",
}
ROW_HIGHLIGHT_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9E2F3")
KEY_HEADER_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
MODULE_ROW_FILLS = (
    PatternFill(fill_type="solid", fgColor="EAF3FF"),
    PatternFill(fill_type="solid", fgColor="EAFBF1"),
    PatternFill(fill_type="solid", fgColor="FFF4E5"),
    PatternFill(fill_type="solid", fgColor="F8ECFF"),
    PatternFill(fill_type="solid", fgColor="FDEBEC"),
    PatternFill(fill_type="solid", fgColor="EEF2FF"),
    PatternFill(fill_type="solid", fgColor="E8F6F3"),
    PatternFill(fill_type="solid", fgColor="FFF9DB"),
)
_NORMAL_DIST = NormalDist()


def _format_duration(seconds: float) -> str:
    total_seconds = max(int(round(seconds)), 0)
    hours, remainder = divmod(total_seconds, 3600)
    minutes, secs = divmod(remainder, 60)
    if hours:
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


def _log(message: str) -> None:
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {message}")


def _count_file_lines(path: Path) -> int:
    newline_count = 0
    ends_with_newline = False
    with path.open("rb") as handle:
        while True:
            chunk = handle.read(1024 * 1024)
            if not chunk:
                break
            newline_count += chunk.count(b"\n")
            ends_with_newline = chunk.endswith(b"\n")
    if newline_count == 0:
        return 0
    return newline_count if ends_with_newline else newline_count + 1


@dataclass(frozen=True)
class TestKey:
    test_number: str
    test_name: str

    @property
    def module(self) -> str:
        return self.test_name[:4].upper()


@dataclass
class TestMeta:
    column_index: int
    test_number: str
    test_name: str
    low: float | None
    high: float | None
    unit: str | None


@dataclass
class InsertionHeader:
    insertion: str
    file_path: Path
    columns: list[str]
    meta_by_key: dict[TestKey, TestMeta]
    excluded_monitoring_keys: list[TestKey]
    all_test_names_by_number: dict[str, str]
    device_row_count: int


def _sanitize_filename(text: str) -> str:
    sanitized = re.sub(r"[^A-Za-z0-9._-]+", "_", text.strip())
    sanitized = re.sub(r"_+", "_", sanitized).strip("._")
    return sanitized or "unnamed"


def _parse_float(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if text == "" or text.lower() == "nan":
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _normalize_wafer(value: Any) -> str:
    text = str(value).strip()
    if text == "" or text.lower() == "nan":
        return ""
    try:
        number = float(text)
    except ValueError:
        return text
    if number.is_integer():
        return str(int(number))
    return text


def _normalize_test_number(value: Any) -> str:
    text = str(value).strip()
    if text == "" or text.lower() == "nan":
        return ""
    number = _parse_float(text)
    if number is not None and float(number).is_integer():
        return str(int(number))
    return text


def _module_from_test_name(test_name: str) -> str:
    clean_name = str(test_name).strip()
    prefix = clean_name[:4].upper()
    return prefix if prefix else "UNKN"


def _normalize_module_list(modules: list[str] | tuple[str, ...] | set[str] | None) -> tuple[str, ...]:
    if not modules:
        return DEFAULT_MODULES
    ordered: list[str] = []
    seen: set[str] = set()
    for module in modules:
        clean = str(module).strip().upper()
        if not clean or clean in seen:
            continue
        seen.add(clean)
        ordered.append(clean)
    return tuple(ordered)


def _normalize_monitoring_patterns(patterns: list[str] | tuple[str, ...] | set[str] | None) -> tuple[str, ...]:
    if not patterns:
        return tuple()
    ordered: list[str] = []
    seen: set[str] = set()
    for pattern in patterns:
        clean = str(pattern).strip()
        if not clean:
            continue
        normalized = clean.casefold()
        if normalized in seen:
            continue
        seen.add(normalized)
        ordered.append(normalized)
    return tuple(ordered)


def _is_monitoring_test(test_name: str, monitoring_patterns: tuple[str, ...]) -> bool:
    if not monitoring_patterns:
        return False
    normalized_name = str(test_name).casefold()
    return any(pattern in normalized_name for pattern in monitoring_patterns)


def _classify_functional_block(module: str, test_name: str) -> str:
    text = str(test_name).upper()
    if module == "DPLL" or any(keyword in text for keyword in ("TIME", "STP", "PAT", "SEQ", "DIG", "ELAPS")):
        return "Digital logic patterns"
    if module in {"TXGE", "TXVC"} or any(keyword in text for keyword in ("PTAT", "BIAS", "VSUP", "VLD", "VMIN", "VMAX", "IDC", "ICC", "CUR", "VREF", "VREG")):
        return "Analog bias circuits"
    if any(keyword in text for keyword in ("LO", "PH", "GAIN", "PWR", "POWER", "FW", "RV", "COR", "ACP", "EVM", "MAG", "ATT")):
        return "Microwave S-parameters"
    return "DC parameters"


def _normalize_pass_fail(value: Any) -> str:
    text = str(value).strip().upper()
    if text.startswith("P"):
        return "P"
    if text.startswith("F"):
        return "F"
    return text


def _prepare_chip_frame(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    for column in ("X", "Y", "SITE_NUM"):
        if column in result.columns:
            result[column] = pd.to_numeric(result[column], errors="coerce")
    if "WAFER" in result.columns:
        result["WAFER"] = result["WAFER"].map(_normalize_wafer)
    if "CHIP_ID" in result.columns:
        result["CHIP_ID"] = result["CHIP_ID"].astype(str).str.strip()
        result.loc[result["CHIP_ID"].str.lower() == "nan", "CHIP_ID"] = ""
    if "PF" in result.columns:
        result["PF"] = result["PF"].map(_normalize_pass_fail)
    if "FIRST_FAIL_TEST" in result.columns:
        result["FIRST_FAIL_TEST"] = result["FIRST_FAIL_TEST"].map(_normalize_test_number)

    wafer = result.get("WAFER", pd.Series("", index=result.index)).fillna("").astype(str)
    x = pd.to_numeric(result.get("X", pd.Series(np.nan, index=result.index)), errors="coerce")
    y = pd.to_numeric(result.get("Y", pd.Series(np.nan, index=result.index)), errors="coerce")
    coord_key = np.where(
        wafer.ne("") & np.isfinite(x) & np.isfinite(y),
        wafer + ":" + x.astype("Int64").astype(str) + ":" + y.astype("Int64").astype(str),
        "",
    )
    chip_id = result.get("CHIP_ID", pd.Series("", index=result.index)).fillna("").astype(str).str.strip()
    result["coord_key"] = coord_key
    result["chip_key"] = np.where(chip_id.ne(""), chip_id, coord_key)
    result = result[result["chip_key"] != ""].copy()
    return result


def _finite_array(values: np.ndarray) -> np.ndarray:
    arr = np.asarray(values, dtype=float)
    return arr[np.isfinite(arr)]


def _paired_finite_arrays(left: np.ndarray, right: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    left_arr = np.asarray(left, dtype=float)
    right_arr = np.asarray(right, dtype=float)
    mask = np.isfinite(left_arr) & np.isfinite(right_arr)
    return left_arr[mask], right_arr[mask]


def _pearson_r2(left: np.ndarray, right: np.ndarray) -> tuple[float, float, int]:
    x, y = _paired_finite_arrays(left, right)
    if x.size < 3:
        return np.nan, np.nan, int(x.size)
    if np.std(x, ddof=1) <= 0 or np.std(y, ddof=1) <= 0:
        return np.nan, np.nan, int(x.size)
    corr = float(np.corrcoef(x, y)[0, 1])
    if not math.isfinite(corr):
        return np.nan, np.nan, int(x.size)
    return corr, float(corr * corr), int(x.size)


def _guardband_shift(base_values: np.ndarray, ambient_values: np.ndarray) -> tuple[float, float, float, int]:
    base, ambient = _paired_finite_arrays(base_values, ambient_values)
    if base.size == 0:
        return np.nan, np.nan, np.nan, 0
    diff = ambient - base
    median_delta = float(np.median(diff))
    sigma_delta = float(np.std(diff, ddof=1)) if diff.size > 1 else 0.0
    return abs(median_delta) + 3.0 * sigma_delta, median_delta, sigma_delta, int(diff.size)


def _guardband_for_sigma(base_values: np.ndarray, ambient_values: np.ndarray, sigma_multiplier: float) -> tuple[float, float, float, int]:
    base, ambient = _paired_finite_arrays(base_values, ambient_values)
    if base.size == 0:
        return np.nan, np.nan, np.nan, 0
    diff = ambient - base
    median_delta = float(np.median(diff))
    sigma_delta = float(np.std(diff, ddof=1)) if diff.size > 1 else 0.0
    return abs(median_delta) + sigma_multiplier * sigma_delta, median_delta, sigma_delta, int(diff.size)


def _tighten_limits(low: float | None, high: float | None, guardband: float) -> tuple[float | None, float | None, bool]:
    if not math.isfinite(guardband):
        return low, high, False
    new_low = low + guardband if low is not None and math.isfinite(low) else None
    new_high = high - guardband if high is not None and math.isfinite(high) else None
    if new_low is not None and new_high is not None and new_low >= new_high:
        return new_low, new_high, False
    return new_low, new_high, True


def _fail_mask_for_values(values: np.ndarray, low: float | None, high: float | None) -> np.ndarray:
    arr = np.asarray(values, dtype=float)
    finite_mask = np.isfinite(arr)
    fail_mask = np.zeros(arr.shape[0], dtype=bool)
    if low is not None and math.isfinite(low):
        fail_mask[finite_mask] |= arr[finite_mask] < float(low)
    if high is not None and math.isfinite(high):
        fail_mask[finite_mask] |= arr[finite_mask] > float(high)
    return fail_mask


def _ppm(count: int, denominator: int) -> float:
    if denominator <= 0:
        return np.nan
    return float(count * 1_000_000.0 / denominator)


def _zero_failure_ucb_ppm(screened_chip_count: int, confidence_level: float) -> float:
    if screened_chip_count <= 0 or not (0.0 < confidence_level < 1.0):
        return np.nan
    alpha = 1.0 - confidence_level
    return float((-math.log(alpha) / screened_chip_count) * 1_000_000.0)


def _closest_limit_margin(value: float | None, low: float | None, high: float | None) -> float | None:
    if value is None or not math.isfinite(value):
        return None
    margins: list[float] = []
    if low is not None and math.isfinite(low):
        margins.append(float(value) - float(low))
    if high is not None and math.isfinite(high):
        margins.append(float(high) - float(value))
    if not margins:
        return None
    return float(min(margins, key=abs))


def _closest_limit_margin_array(values: np.ndarray, low: float | None, high: float | None) -> np.ndarray:
    margins: list[float] = []
    for value in np.asarray(values, dtype=float):
        margin = _closest_limit_margin(value, low, high)
        if margin is not None and math.isfinite(margin):
            margins.append(float(margin))
    return np.asarray(margins, dtype=float)


def _normalized_transfer_metrics(base_values: np.ndarray, ambient_values: np.ndarray, ambient_low: float | None, ambient_high: float | None) -> dict[str, Any]:
    paired_base, paired_ambient = _paired_finite_arrays(base_values, ambient_values)
    if paired_base.size == 0:
        return {"margin_ref": np.nan, "normalized_abs_median_delta": np.nan, "normalized_sigma_delta": np.nan}

    ambient_margins = _closest_limit_margin_array(paired_ambient, ambient_low, ambient_high)
    positive_ambient_margins = ambient_margins[ambient_margins > 0.0]
    if positive_ambient_margins.size == 0:
        return {"margin_ref": np.nan, "normalized_abs_median_delta": np.nan, "normalized_sigma_delta": np.nan}

    diff = paired_ambient - paired_base
    median_delta = float(np.median(diff))
    sigma_delta = float(np.std(diff, ddof=1)) if diff.size > 1 else 0.0
    margin_ref = float(np.quantile(positive_ambient_margins, 0.01))
    if not math.isfinite(margin_ref) or margin_ref <= 0.0:
        return {"margin_ref": margin_ref, "normalized_abs_median_delta": np.nan, "normalized_sigma_delta": np.nan}

    return {
        "margin_ref": margin_ref,
        "normalized_abs_median_delta": abs(median_delta) / margin_ref,
        "normalized_sigma_delta": sigma_delta / margin_ref,
    }


def _zero_escape_confidence_tier(analysis_path: str, best_r2: float, unique_escape_chips: int, zero_escape_ucb_95_ppm: float, worst_normalized_transfer: float) -> str:
    if unique_escape_chips > 0:
        return "Escape Observed"
    if not math.isfinite(zero_escape_ucb_95_ppm):
        return "Insufficient Data"
    if analysis_path == "Parametric Guardband":
        if math.isfinite(best_r2) and best_r2 >= 0.98 and math.isfinite(worst_normalized_transfer) and worst_normalized_transfer <= 0.25 and zero_escape_ucb_95_ppm <= 150.0:
            return "Strong"
        if math.isfinite(best_r2) and best_r2 >= 0.95 and math.isfinite(worst_normalized_transfer) and worst_normalized_transfer <= 0.50 and zero_escape_ucb_95_ppm <= 150.0:
            return "Moderate"
        return "Limited"
    if zero_escape_ucb_95_ppm <= 150.0:
        return "Moderate"
    return "Limited"


def _zero_escape_recommended_action(analysis_path: str, confidence_tier: str) -> str:
    if analysis_path == "Go-NoGo Screening":
        if confidence_tier == "Moderate":
            return "No observed S3 escape. Productive go/no-go screening looks stable, but review zero-failure UCB before removing S3."
        return "No observed S3 escape, but confidence is limited. Keep S3 or gather more volume."
    if analysis_path == "One-Sided Screening":
        if confidence_tier == "Moderate":
            return "No observed S3 escape. One-sided screening looks stable, but review zero-failure UCB before removing S3."
        return "No observed S3 escape, but confidence is limited. Keep S3 or gather more volume."
    if confidence_tier == "Strong":
        return "No observed S3 escape. Zero-failure UCB and transfer stability support S3 removal review."
    if confidence_tier == "Moderate":
        return "No observed S3 escape. Transfer looks acceptable, but review zero-failure UCB before removing S3."
    return "No observed S3 escape, but transfer confidence is limited. Keep S3 or gather more volume."


def _limit_mode(low: float | None, high: float | None) -> str:
    low_valid = low is not None and math.isfinite(low)
    high_valid = high is not None and math.isfinite(high)
    if low_valid and high_valid:
        if math.isclose(float(low), float(high), rel_tol=0.0, abs_tol=1e-12):
            return "Go-NoGo"
        return "Two-Sided"
    if low_valid:
        return "Low-Only"
    if high_valid:
        return "High-Only"
    return "No-Limits"


def _one_sided_screen_tighten(base_values: np.ndarray, low: float | None, high: float | None, s3_fail_mask: np.ndarray, screened_chip_count: int) -> dict[str, Any]:
    result: dict[str, Any] = {
        "limit_mode": _limit_mode(low, high),
        "tighten_amount": np.nan,
        "new_low": low,
        "new_high": high,
        "valid": False,
        "capture_all": False,
        "recaptured": 0,
        "residual": int(np.count_nonzero(s3_fail_mask)),
        "residual_ppm": _ppm(int(np.count_nonzero(s3_fail_mask)), screened_chip_count),
        "overkill": 0,
        "overkill_ppm": np.nan,
    }

    mode = result["limit_mode"]
    if mode not in {"Low-Only", "High-Only"}:
        return result

    base_arr = np.asarray(base_values, dtype=float)
    finite_mask = np.isfinite(base_arr)
    escape_base = base_arr[finite_mask & s3_fail_mask]
    if escape_base.size == 0:
        result.update({"valid": True, "capture_all": True, "residual": 0, "residual_ppm": 0.0, "overkill_ppm": 0.0, "tighten_amount": 0.0})
        return result

    if mode == "Low-Only":
        if low is None or not math.isfinite(low):
            return result
        new_low = float(np.nextafter(np.max(escape_base), np.inf))
        new_high = high
        tighten_amount = new_low - float(low)
    else:
        if high is None or not math.isfinite(high):
            return result
        new_low = low
        new_high = float(np.nextafter(np.min(escape_base), -np.inf))
        tighten_amount = float(high) - new_high

    fail_mask = _fail_mask_for_values(base_arr, new_low, new_high)
    recaptured = int(np.count_nonzero(s3_fail_mask & fail_mask))
    residual = int(np.count_nonzero(s3_fail_mask & ~fail_mask))
    overkill = int(np.count_nonzero(~s3_fail_mask & fail_mask))
    result.update(
        {
            "tighten_amount": tighten_amount,
            "new_low": new_low,
            "new_high": new_high,
            "valid": True,
            "capture_all": residual == 0,
            "recaptured": recaptured,
            "residual": residual,
            "residual_ppm": _ppm(residual, screened_chip_count),
            "overkill": overkill,
            "overkill_ppm": _ppm(overkill, screened_chip_count),
        }
    )
    return result


def _correlation_tier(best_r2: float) -> str:
    return "High" if math.isfinite(best_r2) and best_r2 >= HIGH_CORRELATION_R2_THRESHOLD else "Low"


def _escape_tier(escape_ppm: float) -> str:
    if not math.isfinite(escape_ppm) or escape_ppm <= 0.0:
        return "Zero"
    if escape_ppm <= LOW_ESCAPE_PPM_THRESHOLD:
        return "Low"
    return "High"


def _recommended_action(correlation_tier: str, escape_ppm: float, residual_guardband_ppm: float) -> str:
    escape_tier = _escape_tier(escape_ppm)
    if correlation_tier == "High" and escape_tier == "Zero":
        return "Remove S3 module immediately."
    if correlation_tier == "High" and escape_tier == "Low" and (not math.isfinite(residual_guardband_ppm) or residual_guardband_ppm <= 0.0):
        return "Guardband S1/S2 limits to cover S3, then remove."
    if correlation_tier == "Low" and escape_tier == "Zero":
        return "Monitor. The parameter drifts unpredictably but has not caused yield loss yet."
    return "Keep S3. The module still carries unique ambient risk."


def _find_sigma_capture(
    base_values: np.ndarray,
    ambient_values: np.ndarray,
    low: float | None,
    high: float | None,
    s3_fail_mask: np.ndarray,
    screened_chip_count: int,
) -> dict[str, Any]:
    best_result: dict[str, Any] = {
        "sigma_multiplier": np.nan,
        "guardband": np.nan,
        "median_delta": np.nan,
        "sigma_delta": np.nan,
        "paired_n": 0,
        "new_low": low,
        "new_high": high,
        "valid": False,
        "capture_all": False,
        "recaptured": 0,
        "residual": int(np.count_nonzero(s3_fail_mask)),
        "residual_ppm": _ppm(int(np.count_nonzero(s3_fail_mask)), screened_chip_count),
        "overkill": 0,
        "overkill_ppm": 0.0,
    }

    s3_fail_count = int(np.count_nonzero(s3_fail_mask))
    if s3_fail_count == 0:
        best_result.update({"valid": True, "capture_all": True, "sigma_multiplier": 0.0, "residual": 0, "residual_ppm": 0.0})
        return best_result

    for sigma_multiplier in SIGMA_SEARCH_VALUES:
        guardband, median_delta, sigma_delta, paired_n = _guardband_for_sigma(base_values, ambient_values, sigma_multiplier)
        new_low, new_high, valid = _tighten_limits(low, high, guardband)
        if not valid:
            continue
        gb_fail_mask = _fail_mask_for_values(base_values, new_low, new_high)
        recaptured = int(np.count_nonzero(s3_fail_mask & gb_fail_mask))
        residual = int(np.count_nonzero(s3_fail_mask & ~gb_fail_mask))
        overkill = int(np.count_nonzero(~s3_fail_mask & gb_fail_mask))
        result = {
            "sigma_multiplier": sigma_multiplier,
            "guardband": guardband,
            "median_delta": median_delta,
            "sigma_delta": sigma_delta,
            "paired_n": paired_n,
            "new_low": new_low,
            "new_high": new_high,
            "valid": True,
            "capture_all": residual == 0,
            "recaptured": recaptured,
            "residual": residual,
            "residual_ppm": _ppm(residual, screened_chip_count),
            "overkill": overkill,
            "overkill_ppm": _ppm(overkill, screened_chip_count),
        }
        best_result = result
        if residual == 0:
            return result

    return best_result


def _blank_adaptive_guardband(low: float | None, high: float | None) -> dict[str, Any]:
    return {
        "sigma_multiplier": np.nan,
        "guardband": np.nan,
        "median_delta": np.nan,
        "sigma_delta": np.nan,
        "paired_n": 0,
        "new_low": low,
        "new_high": high,
        "valid": False,
        "capture_all": False,
        "recaptured": 0,
        "residual": np.nan,
        "residual_ppm": np.nan,
        "overkill": 0,
        "overkill_ppm": np.nan,
    }


def _discover_input_files() -> dict[str, Path]:
    mapping: dict[str, Path] = {}
    for path in sorted(INPUT_DIR.glob("*.csv")):
        name = path.name.upper()
        if "_S11P_" in name:
            mapping["S1"] = path
        elif "_S21P_" in name:
            mapping["S2"] = path
        elif "_S31P_" in name:
            mapping["S3"] = path
    missing = [item for item in INSERTION_ORDER if item not in mapping]
    if missing:
        raise SystemExit(f"Missing expected input file(s) for: {', '.join(missing)}")
    return mapping


def _read_header(path: Path, insertion: str, selected_modules: set[str], monitoring_patterns: tuple[str, ...] = ()) -> InsertionHeader:
    _log(f"[{insertion}] reading header from {path.name}")
    with path.open("r", encoding=ENCODING, newline="") as handle:
        reader = csv.reader(handle, delimiter=DELIMITER)
        rows = [next(reader) for _ in range(HEADER_ROWS)]

    columns = [str(item).strip() for item in rows[0]]
    test_names = rows[1]
    lows = rows[2]
    highs = rows[3]
    units = rows[4]
    all_test_names_by_number: dict[str, str] = {}

    for col_name, test_name in zip(columns, test_names, strict=False):
        test_number = _normalize_test_number(col_name)
        if test_number:
            all_test_names_by_number[test_number] = str(test_name).strip()

    meta_by_key: dict[TestKey, TestMeta] = {}
    excluded_monitoring_keys: list[TestKey] = []
    for idx, (col_name, test_name) in enumerate(zip(columns, test_names, strict=False)):
        if not str(col_name).isdigit():
            continue
        clean_name = str(test_name).strip()
        if _module_from_test_name(clean_name) not in selected_modules:
            continue
        key = TestKey(test_number=str(col_name), test_name=clean_name)
        if _is_monitoring_test(clean_name, monitoring_patterns):
            excluded_monitoring_keys.append(key)
            continue
        meta_by_key[key] = TestMeta(
            column_index=idx,
            test_number=str(col_name),
            test_name=clean_name,
            low=_parse_float(lows[idx] if idx < len(lows) else None),
            high=_parse_float(highs[idx] if idx < len(highs) else None),
            unit=str(units[idx]).strip() or None,
        )

    total_rows = _count_file_lines(path)

    return InsertionHeader(
        insertion=insertion,
        file_path=path,
        columns=columns,
        meta_by_key=meta_by_key,
        excluded_monitoring_keys=excluded_monitoring_keys,
        all_test_names_by_number=all_test_names_by_number,
        device_row_count=max(total_rows - DATA_START_ROW, 0),
    )


def _build_test_sets(headers: dict[str, InsertionHeader], monitoring_patterns: tuple[str, ...]) -> tuple[list[TestKey], list[dict[str, str]]]:
    common_keys = set.intersection(*(set(header.meta_by_key) for header in headers.values()))
    common_tests = sorted(common_keys, key=lambda item: (item.module, int(item.test_number), item.test_name))

    all_keys = set.union(*(set(header.meta_by_key) for header in headers.values()))
    excluded_monitoring_keys = set.union(*(set(header.excluded_monitoring_keys) for header in headers.values())) if headers else set()
    skipped_rows: list[dict[str, str]] = []
    for key in sorted(excluded_monitoring_keys, key=lambda item: (item.module, int(item.test_number), item.test_name)):
        skipped_rows.append(
            {
                "Module": key.module,
                "Test Number": key.test_number,
                "Test Name": key.test_name,
                "Present S1": "Y" if key in headers["S1"].excluded_monitoring_keys or key in headers["S1"].meta_by_key else "N",
                "Present S2": "Y" if key in headers["S2"].excluded_monitoring_keys or key in headers["S2"].meta_by_key else "N",
                "Present S3": "Y" if key in headers["S3"].excluded_monitoring_keys or key in headers["S3"].meta_by_key else "N",
                "Reason": "Monitoring test excluded by name pattern.",
            }
        )
    for key in sorted(all_keys - common_keys, key=lambda item: (item.module, int(item.test_number), item.test_name)):
        if key in excluded_monitoring_keys:
            continue
        skipped_rows.append(
            {
                "Module": key.module,
                "Test Number": key.test_number,
                "Test Name": key.test_name,
                "Present S1": "Y" if key in headers["S1"].meta_by_key else "N",
                "Present S2": "Y" if key in headers["S2"].meta_by_key else "N",
                "Present S3": "Y" if key in headers["S3"].meta_by_key else "N",
                "Reason": "Missing from at least one insertion, so no 3-way plot was generated.",
            }
        )
    return common_tests, skipped_rows


def _module_tests(common_tests: list[TestKey], ordered_modules: tuple[str, ...]) -> dict[str, list[TestKey]]:
    grouped: dict[str, list[TestKey]] = {module: [] for module in ordered_modules}
    for key in common_tests:
        grouped.setdefault(key.module, []).append(key)
    return {module: tests for module, tests in grouped.items() if tests}


def _build_read_columns(header: InsertionHeader, tests: list[TestKey]) -> tuple[list[int], list[str], dict[str, TestMeta]]:
    usecols: list[int] = []
    names: list[str] = []
    meta_lookup: dict[str, TestMeta] = {}

    for meta_name in META_COLS:
        try:
            idx = header.columns.index(meta_name)
        except ValueError:
            continue
        usecols.append(idx)
        names.append(header.columns[idx])

    for test in tests:
        meta = header.meta_by_key[test]
        usecols.append(meta.column_index)
        names.append(meta.test_number)
        meta_lookup[meta.test_number] = meta

    return usecols, names, meta_lookup


def _read_module_frame(header: InsertionHeader, tests: list[TestKey]) -> tuple[pd.DataFrame, dict[str, TestMeta]]:
    _log(f"[{header.insertion}] loading {len(tests)} selected test column(s)")
    usecols, _, meta_lookup = _build_read_columns(header, tests)
    df = pd.read_csv(
        header.file_path,
        sep=DELIMITER,
        skiprows=DATA_START_ROW,
        header=None,
        names=header.columns,
        usecols=usecols,
        encoding=ENCODING,
        low_memory=False,
    )
    df = _prepare_chip_frame(df)
    for test_number in meta_lookup:
        df[test_number] = pd.to_numeric(df[test_number], errors="coerce")
    return df.drop_duplicates(subset=["chip_key"]), meta_lookup


def _finite_values(series: pd.Series) -> np.ndarray:
    values = pd.to_numeric(series, errors="coerce").to_numpy(dtype=float)
    return values[np.isfinite(values)]


def _cpk(values: np.ndarray, low: float | None, high: float | None) -> float | None:
    if values.size < 2:
        return None
    sigma = float(np.std(values, ddof=1))
    if not np.isfinite(sigma) or sigma <= 0.0:
        return None
    mean = float(np.mean(values))
    candidates: list[float] = []
    if low is not None and math.isfinite(low):
        candidates.append((mean - low) / (3.0 * sigma))
    if high is not None and math.isfinite(high):
        candidates.append((high - mean) / (3.0 * sigma))
    if not candidates:
        return None
    return float(min(candidates))


def _compute_stats(values: np.ndarray, low: float | None, high: float | None) -> dict[str, Any]:
    result: dict[str, Any] = {
        "n": int(values.size),
        "fail_count": 0,
        "fail_pct": np.nan,
        "yield_pct": np.nan,
        "mean": np.nan,
        "median": np.nan,
        "std": np.nan,
        "min": np.nan,
        "max": np.nan,
        "p01": np.nan,
        "p99": np.nan,
        "cpk": np.nan,
    }
    if values.size == 0:
        return result

    fail_mask = np.zeros(values.size, dtype=bool)
    if low is not None and math.isfinite(low):
        fail_mask |= values < float(low)
    if high is not None and math.isfinite(high):
        fail_mask |= values > float(high)

    fail_count = int(np.count_nonzero(fail_mask))
    result.update(
        {
            "fail_count": fail_count,
            "fail_pct": (100.0 * fail_count / values.size),
            "yield_pct": (100.0 * (values.size - fail_count) / values.size),
            "mean": float(np.mean(values)),
            "median": float(np.median(values)),
            "std": float(np.std(values, ddof=1)) if values.size > 1 else np.nan,
            "min": float(np.min(values)),
            "max": float(np.max(values)),
            "p01": float(np.percentile(values, 1)),
            "p99": float(np.percentile(values, 99)),
            "cpk": _cpk(values, low, high),
        }
    )
    return result


def _fmt_number(value: Any, digits: int = 4) -> str:
    if value is None:
        return "N/A"
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if not math.isfinite(number):
        return "N/A"
    return f"{number:.{digits}g}"


def _probability_axis_forward(percent_values: Any) -> np.ndarray:
    arr = np.asarray(percent_values, dtype=float)
    clipped = np.clip(arr / 100.0, 1e-8, 1.0 - 1e-8)
    return np.asarray([_NORMAL_DIST.inv_cdf(float(item)) for item in clipped.ravel()], dtype=float).reshape(arr.shape)


def _probability_axis_inverse(z_values: Any) -> np.ndarray:
    arr = np.asarray(z_values, dtype=float)
    return np.asarray([_NORMAL_DIST.cdf(float(item)) * 100.0 for item in arr.ravel()], dtype=float).reshape(arr.shape)


def _apply_probability_axis(ax: plt.Axes) -> None:
    tick_labels = [f"{tick:g}" for tick in PROBABILITY_TICKS_PCT]
    ax.set_yscale("function", functions=(_probability_axis_forward, _probability_axis_inverse))
    ax.set_ylim(PROBABILITY_TICKS_PCT[0], PROBABILITY_TICKS_PCT[-1])
    ax.set_yticks(PROBABILITY_TICKS_PCT)
    ax.set_yticklabels(tick_labels)
    ax.set_ylabel("CDF (%)")


def _resolve_xlim(values_by_insertion: dict[str, np.ndarray], limits: list[float]) -> tuple[float, float] | None:
    finite_arrays = [arr[np.isfinite(arr)] for arr in values_by_insertion.values() if arr.size]
    if not finite_arrays and not limits:
        return None
    if finite_arrays:
        merged = np.concatenate(finite_arrays)
        left = float(np.percentile(merged, 0.5))
        right = float(np.percentile(merged, 99.5))
        if merged.size > 20:
            q25, q75 = np.percentile(merged, [25, 75])
            iqr = float(q75 - q25)
            if iqr > 0:
                left = min(left, float(q25 - 2.0 * iqr))
                right = max(right, float(q75 + 2.0 * iqr))
        else:
            left = float(np.min(merged))
            right = float(np.max(merged))
    else:
        left = min(limits)
        right = max(limits)

    if limits:
        left = min([left, *limits])
        right = max([right, *limits])
    if not math.isfinite(left) or not math.isfinite(right):
        return None
    if right <= left:
        pad = max(abs(left) * 0.05, 1.0)
        return left - pad, right + pad
    pad = max((right - left) * 0.06, 1e-9)
    return left - pad, right + pad


def _build_stats_box(stats_by_insertion: dict[str, dict[str, Any]], limits_by_insertion: dict[str, tuple[float | None, float | None]]) -> str:
    lines: list[str] = []
    for insertion in INSERTION_ORDER:
        stats = stats_by_insertion[insertion]
        low, high = limits_by_insertion[insertion]
        lines.append(
            " | ".join(
                [
                    INSERTION_LABELS[insertion],
                    f"n={stats['n']}",
                    f"fails={stats['fail_count']}",
                    f"LTL={_fmt_number(low)}",
                    f"UTL={_fmt_number(high)}",
                    f"mean={_fmt_number(stats['mean'])}",
                    f"median={_fmt_number(stats['median'])}",
                    f"std={_fmt_number(stats['std'])}",
                    f"Cpk={_fmt_number(stats['cpk'])}",
                ]
            )
        )
    return "\n".join(lines)


def _maybe_plot_spec_band(ax: plt.Axes, insertion: str, low: float | None, high: float | None) -> None:
    if low is None or high is None:
        return
    if not math.isfinite(low) or not math.isfinite(high):
        return
    left = float(min(low, high))
    right = float(max(low, high))
    ax.axvspan(left, right, color=INSERTION_COLORS[insertion], alpha=0.05, zorder=0)


def _downsample_cdf_points(values: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    sorted_values = np.sort(values)
    if sorted_values.size == 0:
        return sorted_values, np.empty(0, dtype=float)
    if sorted_values.size <= MAX_CDF_PLOT_POINTS:
        y = 100.0 * np.arange(1, sorted_values.size + 1) / sorted_values.size
        return sorted_values, y

    indices = np.linspace(0, sorted_values.size - 1, MAX_CDF_PLOT_POINTS, dtype=int)
    indices = np.unique(indices)
    return sorted_values[indices], 100.0 * (indices + 1) / sorted_values.size


def _plot_cdf(ax: plt.Axes, values_by_insertion: dict[str, np.ndarray], limits_by_insertion: dict[str, tuple[float | None, float | None]]) -> None:
    all_limits: list[float] = []

    for insertion in INSERTION_ORDER:
        plot_values, plot_y = _downsample_cdf_points(values_by_insertion[insertion])
        if plot_values.size == 0:
            continue
        ax.scatter(
            plot_values,
            plot_y,
            color=INSERTION_COLORS[insertion],
            s=6,
            alpha=0.75,
            edgecolors="none",
            label=f"{INSERTION_LABELS[insertion]} (n={values_by_insertion[insertion].size})",
            rasterized=True,
        )

    for insertion in INSERTION_ORDER:
        low, high = limits_by_insertion[insertion]
        _maybe_plot_spec_band(ax, insertion, low, high)
        for label, value in ((f"LTL {insertion}", low), (f"UTL {insertion}", high)):
            if value is None or not math.isfinite(value):
                continue
            all_limits.append(float(value))
            ax.axvline(
                float(value),
                color=INSERTION_COLORS[insertion],
                linestyle=INSERTION_LIMIT_STYLES[insertion],
                linewidth=1.15,
                alpha=0.95,
                label=f"{label}={_fmt_number(value)}",
            )

    xlim = _resolve_xlim(values_by_insertion, all_limits)
    if xlim is not None:
        ax.set_xlim(*xlim)
    _apply_probability_axis(ax)
    ax.set_xlabel("Test value")
    ax.grid(True, alpha=0.25)


def _plot_test(
    test: TestKey,
    module_dir: Path,
    values_by_insertion: dict[str, np.ndarray],
    stats_by_insertion: dict[str, dict[str, Any]],
    limits_by_insertion: dict[str, tuple[float | None, float | None]],
    ) -> Path:
    s3_fail_count = int(stats_by_insertion["S3"]["fail_count"])
    fig, ax_cdf = plt.subplots(figsize=(10, 5.2), dpi=105)

    _plot_cdf(ax_cdf, values_by_insertion, limits_by_insertion)
    ax_cdf.set_title(
        f"{test.module} | {test.test_number} | {test.test_name}\nS3 failing chips: {s3_fail_count}",
        fontsize=12,
    )
    stats_box = _build_stats_box(stats_by_insertion, limits_by_insertion)
    ax_cdf.text(
        0.01,
        0.01,
        stats_box,
        transform=ax_cdf.transAxes,
        ha="left",
        va="bottom",
        fontsize=8,
        family="monospace",
        bbox={"facecolor": "white", "alpha": 0.92, "edgecolor": "#BFBFBF", "boxstyle": "round,pad=0.4"},
    )
    ax_cdf.legend(loc="upper left", fontsize=8, framealpha=0.9)

    fig.subplots_adjust(left=0.08, right=0.985, top=0.9, bottom=0.1)
    output_path = module_dir / f"{test.test_number}_{_sanitize_filename(test.test_name)}.png"
    fig.savefig(output_path)
    plt.close(fig)
    return output_path


def _read_chip_meta_frame(header: InsertionHeader) -> pd.DataFrame:
    _log(f"[{header.insertion}] loading chip-level metadata")
    wanted_cols = [column for column in ("WAFER", "X", "Y", "CHIP_ID", "PF", "HBIN", "SBIN", "FIRST_FAIL_TEST") if column in header.columns]
    selected_indices = [header.columns.index(column) for column in wanted_cols]
    records: list[dict[str, Any]] = []

    with header.file_path.open("r", encoding=ENCODING, newline="") as handle:
        reader = csv.reader(handle, delimiter=DELIMITER)
        for _ in range(DATA_START_ROW):
            next(reader, None)
        for row in reader:
            record: dict[str, Any] = {}
            for column, idx in zip(wanted_cols, selected_indices, strict=False):
                record[column] = row[idx] if idx < len(row) else None
            records.append(record)

    df = pd.DataFrame.from_records(records, columns=wanted_cols)
    df = _prepare_chip_frame(df)
    for column in wanted_cols:
        if column not in df.columns:
            df[column] = np.nan
    return df[["chip_key", "coord_key", "WAFER", "X", "Y", "CHIP_ID", "PF", "HBIN", "SBIN", "FIRST_FAIL_TEST"]].drop_duplicates(subset=["chip_key"])


def _build_master_chip_frame(headers: dict[str, InsertionHeader], selected_modules: set[str]) -> tuple[pd.DataFrame, dict[str, pd.DataFrame]]:
    raw_frames = {insertion: _read_chip_meta_frame(headers[insertion]) for insertion in INSERTION_ORDER}
    master = raw_frames["S1"].rename(
        columns={
            "coord_key": "coord_key_S1",
            "PF": "PF_S1",
            "HBIN": "HBIN_S1",
            "SBIN": "SBIN_S1",
            "FIRST_FAIL_TEST": "FIRST_FAIL_TEST_S1",
        }
    )
    for insertion in ("S2", "S3"):
        renamed = raw_frames[insertion].rename(
            columns={
                "coord_key": f"coord_key_{insertion}",
                "PF": f"PF_{insertion}",
                "HBIN": f"HBIN_{insertion}",
                "SBIN": f"SBIN_{insertion}",
                "FIRST_FAIL_TEST": f"FIRST_FAIL_TEST_{insertion}",
            }
        )
        keep_cols = ["chip_key", f"coord_key_{insertion}", f"PF_{insertion}", f"HBIN_{insertion}", f"SBIN_{insertion}", f"FIRST_FAIL_TEST_{insertion}"]
        master = master.merge(renamed[keep_cols], on="chip_key", how="inner")

    master["COORD_MATCH_ALL"] = (master["coord_key_S1"] == master["coord_key_S2"]) & (master["coord_key_S1"] == master["coord_key_S3"])
    master["IS_SCREENED"] = (master["PF_S1"] == "P") & (master["PF_S2"] == "P")
    master["S3_LOT_FAIL"] = master["PF_S3"] != "P"
    master["FIRST_FAIL_TEST_S3"] = master["FIRST_FAIL_TEST_S3"].map(_normalize_test_number)
    master["S3_FIRST_FAIL_NAME"] = master["FIRST_FAIL_TEST_S3"].map(headers["S3"].all_test_names_by_number).fillna("UNKNOWN")
    master["S3_FIRST_FAIL_MODULE"] = master["S3_FIRST_FAIL_NAME"].map(_module_from_test_name)
    master["S3_FIRST_FAIL_BLOCK"] = [
        _classify_functional_block(module, test_name)
        for module, test_name in zip(master["S3_FIRST_FAIL_MODULE"], master["S3_FIRST_FAIL_NAME"], strict=False)
    ]
    master["S3_FIRST_FAIL_IN_SCOPE"] = master["S3_FIRST_FAIL_MODULE"].isin(selected_modules)
    return master, raw_frames


def _build_alignment_summary_sheet(headers: dict[str, InsertionHeader], raw_frames: dict[str, pd.DataFrame], master_chip_df: pd.DataFrame, selected_modules: tuple[str, ...]) -> pd.DataFrame:
    screened_chips = int(master_chip_df["IS_SCREENED"].sum())
    ambient_lot_fails = int((master_chip_df["IS_SCREENED"] & master_chip_df["S3_LOT_FAIL"]).sum())
    in_scope_ambient_first_fails = int((master_chip_df["IS_SCREENED"] & master_chip_df["S3_LOT_FAIL"] & master_chip_df["S3_FIRST_FAIL_IN_SCOPE"]).sum())

    rows: list[dict[str, Any]] = []
    for insertion in INSERTION_ORDER:
        rows.extend(
            [
                {"Metric": f"{insertion} file", "Value": headers[insertion].file_path.name},
                {"Metric": f"{insertion} device rows in file", "Value": headers[insertion].device_row_count},
                {"Metric": f"{insertion} usable chip rows", "Value": len(raw_frames[insertion])},
                {"Metric": f"{insertion} selected-module tests found", "Value": len(headers[insertion].meta_by_key)},
            ]
        )

    rows.extend(
        [
            {"Metric": "Selected modules", "Value": ", ".join(selected_modules)},
            {"Metric": "Matched S1/S2/S3 chips", "Value": len(master_chip_df)},
            {"Metric": "Screened chips passing S1 and S2", "Value": screened_chips},
            {"Metric": "Screened chips failing S3 lot result", "Value": ambient_lot_fails},
            {"Metric": "Screened S3 lot fail PPM", "Value": _ppm(ambient_lot_fails, screened_chips)},
            {"Metric": "Screened S3 lot fails with selected-module first fail", "Value": in_scope_ambient_first_fails},
            {"Metric": "Screened S3 lot fail PPM with selected-module first fail", "Value": _ppm(in_scope_ambient_first_fails, screened_chips)},
            {"Metric": "Coordinate mismatch count across insertions", "Value": int((~master_chip_df["COORD_MATCH_ALL"]).sum())},
            {"Metric": "High-correlation R^2 threshold", "Value": HIGH_CORRELATION_R2_THRESHOLD},
            {"Metric": "Low escape threshold (PPM)", "Value": LOW_ESCAPE_PPM_THRESHOLD},
        ]
    )
    return pd.DataFrame(rows)


def _build_escape_pareto_sheets(master_chip_df: pd.DataFrame, screened_chip_count: int) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    escapes = master_chip_df[master_chip_df["IS_SCREENED"] & master_chip_df["S3_LOT_FAIL"]].copy()
    if escapes.empty:
        empty = pd.DataFrame()
        return empty, empty, empty

    module_pareto = (
        escapes.groupby(["S3_FIRST_FAIL_MODULE", "S3_FIRST_FAIL_IN_SCOPE"], dropna=False)
        .agg(Escape_Chips=("chip_key", "nunique"))
        .reset_index()
        .sort_values(["Escape_Chips", "S3_FIRST_FAIL_MODULE"], ascending=[False, True])
    )
    module_pareto["Escape PPM"] = module_pareto["Escape_Chips"].map(lambda value: _ppm(int(value), screened_chip_count))
    module_pareto = module_pareto.rename(
        columns={
            "S3_FIRST_FAIL_MODULE": "First Fail Module",
            "S3_FIRST_FAIL_IN_SCOPE": "Is Selected Module",
        }
    )
    module_pareto["Is Selected Module"] = module_pareto["Is Selected Module"].map(lambda flag: "Y" if flag else "N")

    test_pareto = (
        escapes.groupby(["FIRST_FAIL_TEST_S3", "S3_FIRST_FAIL_NAME", "S3_FIRST_FAIL_MODULE", "S3_FIRST_FAIL_IN_SCOPE"], dropna=False)
        .agg(Escape_Chips=("chip_key", "nunique"))
        .reset_index()
        .sort_values(["Escape_Chips", "FIRST_FAIL_TEST_S3"], ascending=[False, True])
    )
    test_pareto["Escape PPM"] = test_pareto["Escape_Chips"].map(lambda value: _ppm(int(value), screened_chip_count))
    test_pareto = test_pareto.rename(
        columns={
            "FIRST_FAIL_TEST_S3": "First Fail Test",
            "S3_FIRST_FAIL_NAME": "First Fail Name",
            "S3_FIRST_FAIL_MODULE": "First Fail Module",
            "S3_FIRST_FAIL_IN_SCOPE": "Is Selected Module",
        }
    )
    test_pareto["Is Selected Module"] = test_pareto["Is Selected Module"].map(lambda flag: "Y" if flag else "N")

    chip_details = escapes[
        [
            "chip_key",
            "WAFER",
            "X",
            "Y",
            "PF_S1",
            "PF_S2",
            "PF_S3",
            "HBIN_S3",
            "SBIN_S3",
            "FIRST_FAIL_TEST_S3",
            "S3_FIRST_FAIL_NAME",
            "S3_FIRST_FAIL_MODULE",
            "S3_FIRST_FAIL_IN_SCOPE",
        ]
    ].copy()
    chip_details = chip_details.rename(
        columns={
            "chip_key": "Chip Key",
            "WAFER": "Wafer",
            "X": "X",
            "Y": "Y",
            "PF_S1": "PF S1",
            "PF_S2": "PF S2",
            "PF_S3": "PF S3",
            "HBIN_S3": "HBIN S3",
            "SBIN_S3": "SBIN S3",
            "FIRST_FAIL_TEST_S3": "First Fail Test",
            "S3_FIRST_FAIL_NAME": "First Fail Name",
            "S3_FIRST_FAIL_MODULE": "First Fail Module",
            "S3_FIRST_FAIL_IN_SCOPE": "Is Selected Module",
        }
    )
    chip_details["Is Selected Module"] = chip_details["Is Selected Module"].map(lambda flag: "Y" if flag else "N")
    return module_pareto, test_pareto, chip_details


def _build_module_measurement_frame(master_chip_df: pd.DataFrame, frames: dict[str, pd.DataFrame], tests: list[TestKey]) -> pd.DataFrame:
    module_frame = master_chip_df[
        [
            "chip_key",
            "WAFER",
            "X",
            "Y",
            "CHIP_ID",
            "IS_SCREENED",
            "S3_LOT_FAIL",
            "FIRST_FAIL_TEST_S3",
            "S3_FIRST_FAIL_NAME",
            "S3_FIRST_FAIL_MODULE",
            "S3_FIRST_FAIL_IN_SCOPE",
        ]
    ].copy()
    for insertion in INSERTION_ORDER:
        rename_map = {test.test_number: f"{test.test_number}_{insertion}" for test in tests}
        selected = frames[insertion][["chip_key", *rename_map.keys()]].rename(columns=rename_map)
        module_frame = module_frame.merge(selected, on="chip_key", how="inner")
    return module_frame


def _select_preferred_guardband(recaptured_s1: int, overkill_s1_ppm: float, r2_s1: float, recaptured_s2: int, overkill_s2_ppm: float, r2_s2: float) -> str:
    candidates = [
        ("S1", recaptured_s1, overkill_s1_ppm if math.isfinite(overkill_s1_ppm) else float("inf"), r2_s1 if math.isfinite(r2_s1) else -1.0),
        ("S2", recaptured_s2, overkill_s2_ppm if math.isfinite(overkill_s2_ppm) else float("inf"), r2_s2 if math.isfinite(r2_s2) else -1.0),
    ]
    candidates.sort(key=lambda item: (-item[1], item[2], -item[3], item[0]))
    if candidates[0][1] <= 0:
        return "None"
    if len(candidates) > 1 and candidates[0][1] == candidates[1][1] and candidates[0][2] == candidates[1][2] and candidates[0][3] == candidates[1][3]:
        return "Either"
    return candidates[0][0]


def _collect_report_row(
    test: TestKey,
    functional_block: str,
    plot_path: Path | None,
    screened_chip_count: int,
    analysis_path: str,
    limit_mode: str,
    stats_by_insertion: dict[str, dict[str, Any]],
    limits_by_insertion: dict[str, tuple[float | None, float | None]],
    unit_by_insertion: dict[str, str | None],
    correlation_by_pair: dict[str, dict[str, Any]],
    transfer_confidence_by_pair: dict[str, dict[str, Any]],
    guardband_by_insertion: dict[str, dict[str, Any]],
    adaptive_guardband_by_insertion: dict[str, dict[str, Any]],
    current_screen_metrics: dict[str, Any],
    one_sided_screen_by_insertion: dict[str, dict[str, Any]],
    escape_metrics: dict[str, Any],
) -> dict[str, Any]:
    best_r2 = max(
        [value["r2"] for value in correlation_by_pair.values() if math.isfinite(value["r2"])],
        default=np.nan,
    )
    correlation_tier = _correlation_tier(best_r2)
    action = _recommended_action(correlation_tier, escape_metrics["unique_escape_ppm"], escape_metrics["residual_escape_ppm_either"])
    zero_escape_ucb_90_ppm = _zero_failure_ucb_ppm(screened_chip_count, 0.90) if escape_metrics["unique_escape_chips"] == 0 else np.nan
    zero_escape_ucb_95_ppm = _zero_failure_ucb_ppm(screened_chip_count, 0.95) if escape_metrics["unique_escape_chips"] == 0 else np.nan
    zero_escape_ucb_99_ppm = _zero_failure_ucb_ppm(screened_chip_count, 0.99) if escape_metrics["unique_escape_chips"] == 0 else np.nan
    worst_normalized_transfer = max(
        [metrics["normalized_abs_median_delta"] + metrics["normalized_sigma_delta"] for metrics in transfer_confidence_by_pair.values() if math.isfinite(metrics["normalized_abs_median_delta"]) and math.isfinite(metrics["normalized_sigma_delta"])],
        default=np.nan,
    )
    zero_escape_confidence_tier = _zero_escape_confidence_tier(analysis_path, best_r2, escape_metrics["unique_escape_chips"], zero_escape_ucb_95_ppm, worst_normalized_transfer)

    adaptive_guardband_used = bool(escape_metrics.get("adaptive_guardband_used", False))
    row: dict[str, Any] = {
        "Analysis Path": analysis_path,
        "Limit Mode": limit_mode,
        "Module": test.module,
        "Test Number": int(test.test_number),
        "Test Name": test.test_name,
        "Plot Path": str(plot_path) if plot_path is not None else "",
        "Screened Chips": screened_chip_count,
        "Has Unique S3 Escape": "Y" if escape_metrics["unique_escape_chips"] > 0 else "N",
        "Unique S3 Escape Chips": escape_metrics["unique_escape_chips"],
        "Unique S3 Escape PPM": escape_metrics["unique_escape_ppm"],
        "Zero-Escape UCB 90% PPM": zero_escape_ucb_90_ppm,
        "Zero-Escape UCB 95% PPM": zero_escape_ucb_95_ppm,
        "Zero-Escape UCB 99% PPM": zero_escape_ucb_99_ppm,
        "Zero-Escape Confidence Tier": zero_escape_confidence_tier,
        "Recaptured by S1 Guardband": escape_metrics["recaptured_s1"],
        "Recaptured by S2 Guardband": escape_metrics["recaptured_s2"],
        "Recaptured by Either Guardband": escape_metrics["recaptured_either"],
        "Residual Escapes After Either": escape_metrics["residual_escape_either"],
        "Residual Escape PPM After Either": escape_metrics["residual_escape_ppm_either"],
        "Virtual Yield Loss S1 GB Chips": escape_metrics["overkill_s1"],
        "Virtual Yield Loss S2 GB Chips": escape_metrics["overkill_s2"],
        "Virtual Yield Loss Either GB Chips": escape_metrics["overkill_either"],
        "Virtual Yield Loss S1 GB PPM": escape_metrics["overkill_s1_ppm"],
        "Virtual Yield Loss S2 GB PPM": escape_metrics["overkill_s2_ppm"],
        "Virtual Yield Loss Either GB PPM": escape_metrics["overkill_either_ppm"],
        "Preferred Guardband Insertion": escape_metrics["preferred_guardband"],
        "Adaptive Sigma S1": adaptive_guardband_by_insertion["S1"]["sigma_multiplier"] if adaptive_guardband_used else np.nan,
        "Adaptive Sigma S2": adaptive_guardband_by_insertion["S2"]["sigma_multiplier"] if adaptive_guardband_used else np.nan,
        "Adaptive Capture All S1": "Y" if adaptive_guardband_used and adaptive_guardband_by_insertion["S1"]["capture_all"] else "",
        "Adaptive Capture All S2": "Y" if adaptive_guardband_used and adaptive_guardband_by_insertion["S2"]["capture_all"] else "",
        "Adaptive GB S1": adaptive_guardband_by_insertion["S1"]["guardband"] if adaptive_guardband_used else np.nan,
        "Adaptive GB S2": adaptive_guardband_by_insertion["S2"]["guardband"] if adaptive_guardband_used else np.nan,
        "Adaptive Yield Loss S1 PPM": adaptive_guardband_by_insertion["S1"]["overkill_ppm"] if adaptive_guardband_used else np.nan,
        "Adaptive Yield Loss S2 PPM": adaptive_guardband_by_insertion["S2"]["overkill_ppm"] if adaptive_guardband_used else np.nan,
        "Current Screen Recaptured by S1": current_screen_metrics["recaptured_s1"],
        "Current Screen Recaptured by S2": current_screen_metrics["recaptured_s2"],
        "Current Screen Recaptured by Either": current_screen_metrics["recaptured_either"],
        "Current Screen Residual After Either": current_screen_metrics["residual_either"],
        "Current Screen Residual PPM After Either": current_screen_metrics["residual_ppm_either"],
        "One-Sided Tighten S1": one_sided_screen_by_insertion["S1"]["tighten_amount"],
        "One-Sided Tighten S2": one_sided_screen_by_insertion["S2"]["tighten_amount"],
        "One-Sided Capture All S1": "Y" if one_sided_screen_by_insertion["S1"]["capture_all"] else "",
        "One-Sided Capture All S2": "Y" if one_sided_screen_by_insertion["S2"]["capture_all"] else "",
        "One-Sided Yield Loss S1 PPM": one_sided_screen_by_insertion["S1"]["overkill_ppm"],
        "One-Sided Yield Loss S2 PPM": one_sided_screen_by_insertion["S2"]["overkill_ppm"],
        "Transfer Margin Ref S3-S1": transfer_confidence_by_pair["S3-S1"]["margin_ref"],
        "Transfer Margin Ref S3-S2": transfer_confidence_by_pair["S3-S2"]["margin_ref"],
        "Normalized Abs Median Delta S3-S1": transfer_confidence_by_pair["S3-S1"]["normalized_abs_median_delta"],
        "Normalized Abs Median Delta S3-S2": transfer_confidence_by_pair["S3-S2"]["normalized_abs_median_delta"],
        "Normalized Sigma Delta S3-S1": transfer_confidence_by_pair["S3-S1"]["normalized_sigma_delta"],
        "Normalized Sigma Delta S3-S2": transfer_confidence_by_pair["S3-S2"]["normalized_sigma_delta"],
        "Worst Normalized Transfer": worst_normalized_transfer,
        "Pearson r S3-S1": correlation_by_pair["S3-S1"]["r"],
        "R2 S3-S1": correlation_by_pair["S3-S1"]["r2"],
        "Paired N S3-S1": correlation_by_pair["S3-S1"]["n"],
        "Pearson r S3-S2": correlation_by_pair["S3-S2"]["r"],
        "R2 S3-S2": correlation_by_pair["S3-S2"]["r2"],
        "Paired N S3-S2": correlation_by_pair["S3-S2"]["n"],
        "Best R2": best_r2,
        "Correlation Tier": correlation_tier,
        "Recommended Action": action,
        "Unit S1": unit_by_insertion["S1"] or "",
        "Unit S2": unit_by_insertion["S2"] or "",
        "Unit S3": unit_by_insertion["S3"] or "",
    }

    for insertion in INSERTION_ORDER:
        low, high = limits_by_insertion[insertion]
        stats = stats_by_insertion[insertion]
        row[f"LTL {insertion}"] = low
        row[f"UTL {insertion}"] = high
        row[f"N {insertion}"] = stats["n"]
        row[f"Fails {insertion}"] = stats["fail_count"]
        row[f"Fail % {insertion}"] = stats["fail_pct"]
        row[f"Yield % {insertion}"] = stats["yield_pct"]
        row[f"Mean {insertion}"] = stats["mean"]
        row[f"Median {insertion}"] = stats["median"]
        row[f"Std {insertion}"] = stats["std"]
        row[f"Min {insertion}"] = stats["min"]
        row[f"Max {insertion}"] = stats["max"]
        row[f"P01 {insertion}"] = stats["p01"]
        row[f"P99 {insertion}"] = stats["p99"]
        row[f"Cpk {insertion}"] = stats["cpk"]

    for insertion in ("S1", "S2"):
        guardband = guardband_by_insertion[insertion]
        row[f"Guardband {insertion}"] = guardband["guardband"]
        row[f"Median Delta S3-{insertion}"] = guardband["median_delta"]
        row[f"Sigma Delta S3-{insertion}"] = guardband["sigma_delta"]
        row[f"Guardband Paired N {insertion}"] = guardband["paired_n"]
        row[f"Guardbanded LTL {insertion}"] = guardband["new_low"]
        row[f"Guardbanded UTL {insertion}"] = guardband["new_high"]
        row[f"Guardband Valid {insertion}"] = "Y" if guardband["valid"] else "N"

    row["Abs Mean Shift S3-S1"] = abs((row["Mean S3"] if pd.notna(row["Mean S3"]) else np.nan) - (row["Mean S1"] if pd.notna(row["Mean S1"]) else np.nan))
    row["Abs Mean Shift S3-S2"] = abs((row["Mean S3"] if pd.notna(row["Mean S3"]) else np.nan) - (row["Mean S2"] if pd.notna(row["Mean S2"]) else np.nan))
    row["Abs Median Shift S3-S1"] = abs((row["Median S3"] if pd.notna(row["Median S3"]) else np.nan) - (row["Median S1"] if pd.notna(row["Median S1"]) else np.nan))
    row["Abs Median Shift S3-S2"] = abs((row["Median S3"] if pd.notna(row["Median S3"]) else np.nan) - (row["Median S2"] if pd.notna(row["Median S2"]) else np.nan))
    if analysis_path == "Go-NoGo Screening":
        if escape_metrics["unique_escape_chips"] <= 0:
            row["Recommended Action"] = _zero_escape_recommended_action(analysis_path, zero_escape_confidence_tier)
        elif current_screen_metrics["residual_either"] <= 0:
            row["Recommended Action"] = "S1/S2 already screen-out the S3 escape for this go/no-go screening test. No extra tightening needed."
        else:
            row["Recommended Action"] = "Keep S3. Productive go/no-go screening does not recapture the ambient escape."
        return row
    if analysis_path == "One-Sided Screening":
        if escape_metrics["unique_escape_chips"] <= 0:
            row["Recommended Action"] = _zero_escape_recommended_action(analysis_path, zero_escape_confidence_tier)
        elif current_screen_metrics["residual_either"] <= 0:
            row["Recommended Action"] = "S1/S2 already screen-out the S3 escape for this one-sided screening test. No extra tightening needed."
        elif escape_metrics["preferred_guardband"] in ("S1", "S2") and one_sided_screen_by_insertion[escape_metrics["preferred_guardband"]]["capture_all"]:
            row["Recommended Action"] = f"One-sided tightening on {escape_metrics['preferred_guardband']} screens-out the S3 escape. Check simulated yield loss in S1/S2."
        else:
            row["Recommended Action"] = "Keep S3. Productive one-sided screening does not recapture the ambient escape."
        return row
    if escape_metrics["unique_escape_chips"] <= 0:
        row["Recommended Action"] = _zero_escape_recommended_action(analysis_path, zero_escape_confidence_tier)
        return row
    if adaptive_guardband_used and escape_metrics["preferred_guardband"] in ("S1", "S2"):
        sigma_value = adaptive_guardband_by_insertion[escape_metrics["preferred_guardband"]]["sigma_multiplier"]
        row["Recommended Action"] = f"{sigma_value:g}-sigma GB on {escape_metrics['preferred_guardband']} used to screen-out escape in S3. Check simulated yield losses in S1/S2."
    return row


def _build_module_decision_matrix(report_df: pd.DataFrame, escape_detail_df: pd.DataFrame, screened_chip_count: int) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    has_escape_details = not escape_detail_df.empty and "Module" in escape_detail_df.columns
    confidence_rank = {"Strong": 0, "Moderate": 1, "Limited": 2, "Insufficient Data": 3, "Escape Observed": 4}
    for module, group in report_df.groupby("Module", sort=True):
        if has_escape_details:
            module_escape_details = escape_detail_df[escape_detail_df["Module"] == module]
        else:
            module_escape_details = pd.DataFrame(columns=["Chip Key", "Recaptured by Either Guardband"])
        unique_escape_chips = int(module_escape_details["Chip Key"].nunique())
        unique_escape_ppm = _ppm(unique_escape_chips, screened_chip_count)
        recaptured_unique = int(module_escape_details.loc[module_escape_details["Recaptured by Either Guardband"] == "Y", "Chip Key"].nunique())
        residual_unique = int(module_escape_details.loc[module_escape_details["Recaptured by Either Guardband"] != "Y", "Chip Key"].nunique())
        residual_ppm = _ppm(residual_unique, screened_chip_count)
        finite_r2 = group["Best R2"].replace([np.inf, -np.inf], np.nan).dropna()
        median_best_r2 = float(finite_r2.median()) if not finite_r2.empty else np.nan
        correlation_tier = _correlation_tier(median_best_r2)
        high_corr_tests = int((group["Best R2"] >= HIGH_CORRELATION_R2_THRESHOLD).sum())
        paired_corr_tests = int(group["Best R2"].notna().sum())
        pct_high_corr = (100.0 * high_corr_tests / paired_corr_tests) if paired_corr_tests else np.nan
        analysis_counts = group["Analysis Path"].value_counts()
        zero_escape_ucb_90_ppm = float(group["Zero-Escape UCB 90% PPM"].dropna().min()) if group["Zero-Escape UCB 90% PPM"].notna().any() else np.nan
        zero_escape_ucb_95_ppm = float(group["Zero-Escape UCB 95% PPM"].dropna().min()) if group["Zero-Escape UCB 95% PPM"].notna().any() else np.nan
        zero_escape_ucb_99_ppm = float(group["Zero-Escape UCB 99% PPM"].dropna().min()) if group["Zero-Escape UCB 99% PPM"].notna().any() else np.nan
        worst_normalized_transfer = float(group["Worst Normalized Transfer"].replace([np.inf, -np.inf], np.nan).max()) if group["Worst Normalized Transfer"].notna().any() else np.nan
        confidence_tiers = [tier for tier in group["Zero-Escape Confidence Tier"].dropna().astype(str) if tier in confidence_rank]
        module_confidence_tier = max(confidence_tiers, key=lambda tier: confidence_rank[tier]) if confidence_tiers else "Insufficient Data"
        preferred_guardband = "None"
        non_none_guardbands = group.loc[group["Preferred Guardband Insertion"].isin(["S1", "S2", "Either"]), "Preferred Guardband Insertion"]
        if not non_none_guardbands.empty:
            preferred_guardband = str(non_none_guardbands.mode().iat[0])
        if unique_escape_chips > 0:
            recommended_action = "Review test-level simulated yield losses case by case."
        else:
            recommended_action = _zero_escape_recommended_action("Parametric Guardband", module_confidence_tier)
        rows.append(
            {
                "Module": module,
                "Comparable Tests": int(len(group)),
                "Screened Chips": screened_chip_count,
                "Unique S3 Escape Chips": unique_escape_chips,
                "Unique S3 Escape PPM": unique_escape_ppm,
                "Zero-Escape UCB 90% PPM": zero_escape_ucb_90_ppm,
                "Zero-Escape UCB 95% PPM": zero_escape_ucb_95_ppm,
                "Zero-Escape UCB 99% PPM": zero_escape_ucb_99_ppm,
                "Recaptured Unique Escapes by Either Guardband": recaptured_unique,
                "Residual Unique Escapes After Either": residual_unique,
                "Residual Unique Escape PPM After Either": residual_ppm,
                "Worst Normalized Transfer": worst_normalized_transfer,
                "Zero-Escape Confidence Tier": module_confidence_tier,
                "Median Best R2": median_best_r2,
                "% Tests High R2": pct_high_corr,
                "High-R2 Tests": high_corr_tests,
                "Tests With Correlation": paired_corr_tests,
                "Parametric Guardband Tests": int(analysis_counts.get("Parametric Guardband", 0)),
                "One-Sided Screening Tests": int(analysis_counts.get("One-Sided Screening", 0)),
                "Go-NoGo Screening Tests": int(analysis_counts.get("Go-NoGo Screening", 0)),
                "Correlation Tier": correlation_tier,
                "Escape Tier": _escape_tier(unique_escape_ppm),
                "Preferred Guardband Insertion": preferred_guardband,
                "Recommended Action": recommended_action,
            }
        )
    return pd.DataFrame(rows)


def _recommended_analyses_sheet() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Analysis": "S3 vs B2 chip-level correlation",
                "Why it matters": "Confirms whether BE ambient can predict FE ambient behavior at chip level for each selected module.",
                "What to look for": "Pearson/Spearman correlation, offset stability, residual sigma, and outlier clusters by module.",
            },
            {
                "Analysis": "S3 fail-recapture by B2",
                "Why it matters": "Checks whether every chip failing at S3 is already caught at B2 or by another productive insertion.",
                "What to look for": "False-negative count, false-negative ppm, and the exact chips escaping B2 per module.",
            },
            {
                "Analysis": "Margin-to-limit distribution transfer",
                "Why it matters": "Shows how close each module sits to its spec in S3 and whether B2 preserves that distance with enough guardband.",
                "What to look for": "Distribution of min(value-LTL, UTL-value), worst tails, and proposed transfer guardbands.",
            },
            {
                "Analysis": "Zero-escape upper confidence bound",
                "Why it matters": "Converts zero observed S3 escapes into a statistical upper bound on the true escape rate for the screened population.",
                "What to look for": "90%, 95%, and 99% zero-failure UCB in ppm together with transfer-stability metrics before claiming removal readiness.",
            },
        ]
    )


def _report_metadata(headers: dict[str, InsertionHeader], common_tests: list[TestKey], skipped_rows: list[dict[str, str]], master_chip_df: pd.DataFrame, selected_modules: tuple[str, ...], skip_plots: bool, monitoring_patterns: tuple[str, ...]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    screened_chips = int(master_chip_df["IS_SCREENED"].sum())
    ambient_lot_fails = int((master_chip_df["IS_SCREENED"] & master_chip_df["S3_LOT_FAIL"]).sum())
    monitoring_skips = sum(1 for row in skipped_rows if row.get("Reason") == "Monitoring test excluded by name pattern.")
    for insertion in INSERTION_ORDER:
        header = headers[insertion]
        rows.append({"Item": f"{insertion} file", "Value": str(header.file_path.name)})
        rows.append({"Item": f"{insertion} device rows", "Value": header.device_row_count})
        rows.append({"Item": f"{insertion} selected-module tests found", "Value": len(header.meta_by_key)})
    rows.extend(
        [
            {"Item": "Selected modules", "Value": ", ".join(selected_modules)},
            {"Item": "Monitoring test-name patterns excluded", "Value": ", ".join(monitoring_patterns) if monitoring_patterns else "None"},
            {"Item": "Monitoring tests excluded", "Value": monitoring_skips},
            {"Item": "Comparable selected-module tests", "Value": len(common_tests)},
            {"Item": "Skipped selected-module tests", "Value": len(skipped_rows)},
            {"Item": "Matched chips across S1/S2/S3", "Value": len(master_chip_df)},
            {"Item": "Screened chips passing S1 and S2", "Value": screened_chips},
            {"Item": "Ambient lot fails after S1/S2 screen", "Value": ambient_lot_fails},
            {"Item": "Ambient lot fail ppm after S1/S2 screen", "Value": _ppm(ambient_lot_fails, screened_chips)},
            {"Item": "High-correlation R2 threshold", "Value": HIGH_CORRELATION_R2_THRESHOLD},
            {"Item": "Low escape threshold (ppm)", "Value": LOW_ESCAPE_PPM_THRESHOLD},
            {"Item": "Zero-escape UCB alpha at 90% confidence", "Value": 0.10},
            {"Item": "Zero-escape UCB alpha at 95% confidence", "Value": 0.05},
            {"Item": "Zero-escape UCB alpha at 99% confidence", "Value": 0.01},
            {"Item": "Plots skipped", "Value": "Y" if skip_plots else "N"},
            {"Item": "Output folder", "Value": str(OUTPUT_DIR)},
        ]
    )
    return pd.DataFrame(rows)


def _analysis_path_metadata(report_df: pd.DataFrame) -> pd.DataFrame:
    if report_df.empty or "Analysis Path" not in report_df.columns:
        return pd.DataFrame(columns=["Item", "Value"])

    counts = report_df["Analysis Path"].value_counts()
    return pd.DataFrame(
        [
            {"Item": "Parametric guardband tests", "Value": int(counts.get("Parametric Guardband", 0))},
            {"Item": "One-sided screening tests", "Value": int(counts.get("One-Sided Screening", 0))},
            {"Item": "Go-no-go screening tests", "Value": int(counts.get("Go-NoGo Screening", 0))},
        ]
    )


def _style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    report_ws = workbook["Module_Test_Report"]

    header_font = Font(bold=True)
    key_headers = {
        "Best R2",
        "Recommended Action",
        "Unique S3 Escape PPM",
        "Zero-Escape UCB 95% PPM",
        "Zero-Escape Confidence Tier",
        "Worst Normalized Transfer",
        "Residual Escape PPM After Either",
        "Preferred Guardband Insertion",
        "Adaptive Sigma S1",
        "Adaptive Sigma S2",
        "Adaptive Yield Loss S1 PPM",
        "Adaptive Yield Loss S2 PPM",
    }
    for worksheet in workbook.worksheets:
        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = HEADER_FILL
        worksheet.freeze_panes = "A2"
        for column_cells in worksheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
            width = min(max(len(item) for item in values) + 2, 42)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width

    for cell in report_ws[1]:
        if cell.value in key_headers:
            cell.fill = KEY_HEADER_FILL

    header_row = {cell.value: idx for idx, cell in enumerate(report_ws[1], start=1)}
    module_col = header_row.get("Module")
    plot_path_col = header_row.get("Plot Path")
    module_fill_map: dict[str, PatternFill] = {}
    for row_idx in range(2, report_ws.max_row + 1):
        if module_col is not None:
            module_name = str(report_ws.cell(row=row_idx, column=module_col).value or "")
            if module_name not in module_fill_map:
                module_fill_map[module_name] = MODULE_ROW_FILLS[len(module_fill_map) % len(MODULE_ROW_FILLS)]
            for col_idx in range(1, report_ws.max_column + 1):
                report_ws.cell(row=row_idx, column=col_idx).fill = module_fill_map[module_name]
        if plot_path_col is not None:
            cell = report_ws.cell(row=row_idx, column=plot_path_col)
            if cell.value:
                try:
                    cell.hyperlink = Path(str(cell.value)).resolve().as_uri()
                except OSError:
                    cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"

    workbook.save(path)


def run(*, selected_modules: tuple[str, ...] | None = None, max_tests_per_module: int | None = None, skip_plots: bool = False, monitoring_patterns: tuple[str, ...] | None = None) -> None:
    run_started = perf_counter()
    selected_modules = _normalize_module_list(selected_modules)
    monitoring_patterns = _normalize_monitoring_patterns(monitoring_patterns)
    selected_module_set = set(selected_modules)
    _log(f"starting analysis for modules: {', '.join(selected_modules)}")
    _log(f"configuration: skip_plots={'Y' if skip_plots else 'N'}, max_tests_per_module={max_tests_per_module if max_tests_per_module is not None else 'ALL'}")
    _log(f"monitoring exclusions: {', '.join(monitoring_patterns) if monitoring_patterns else 'none'}")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    _log("discovering insertion input files")
    file_map = _discover_input_files()
    _log("reading insertion headers")
    headers = {insertion: _read_header(path, insertion, selected_module_set, monitoring_patterns) for insertion, path in file_map.items()}
    _log("building matched S1/S2/S3 chip master")
    master_chip_df, raw_chip_frames = _build_master_chip_frame(headers, selected_module_set)
    _log("building comparable selected-module test set")
    common_tests, skipped_rows = _build_test_sets(headers, monitoring_patterns)
    tests_by_module = _module_tests(common_tests, selected_modules)
    if max_tests_per_module is not None:
        tests_by_module = {module: tests[:max_tests_per_module] for module, tests in tests_by_module.items()}

    master_chip_df = master_chip_df.copy()
    screened_chip_count = int(master_chip_df["IS_SCREENED"].sum())
    report_rows: list[dict[str, Any]] = []
    escape_detail_rows: list[dict[str, Any]] = []
    total_tests_to_process = sum(len(tests) for tests in tests_by_module.values())
    processed_tests = 0
    _log(f"prepared {total_tests_to_process} comparable test(s) across {len(tests_by_module)} module(s)")

    for module, tests in tests_by_module.items():
        module_started = perf_counter()
        _log(f"[{module}] starting module with {len(tests)} test(s)")
        module_dir = OUTPUT_DIR / module if not skip_plots else None
        if module_dir is not None:
            module_dir.mkdir(parents=True, exist_ok=True)

        frames: dict[str, pd.DataFrame] = {}
        metas: dict[str, dict[str, TestMeta]] = {}
        for insertion in INSERTION_ORDER:
            frame, meta_lookup = _read_module_frame(headers[insertion], tests)
            frames[insertion] = frame
            metas[insertion] = meta_lookup

        _log(f"[{module}] aligning screened chip measurements")
        module_frame = _build_module_measurement_frame(master_chip_df, frames, tests)
        module_screened = module_frame[module_frame["IS_SCREENED"]].reset_index(drop=True)

        for test in tests:
            functional_block = _classify_functional_block(test.module, test.test_name)
            values_by_insertion: dict[str, np.ndarray] = {}
            stats_by_insertion: dict[str, dict[str, Any]] = {}
            limits_by_insertion: dict[str, tuple[float | None, float | None]] = {}
            unit_by_insertion: dict[str, str | None] = {}

            series_by_insertion: dict[str, np.ndarray] = {}
            for insertion in INSERTION_ORDER:
                meta = metas[insertion][test.test_number]
                values = pd.to_numeric(module_screened[f"{test.test_number}_{insertion}"], errors="coerce").to_numpy(dtype=float)
                series_by_insertion[insertion] = values
                finite_values = _finite_array(values)
                values_by_insertion[insertion] = finite_values
                stats_by_insertion[insertion] = _compute_stats(finite_values, meta.low, meta.high)
                limits_by_insertion[insertion] = (meta.low, meta.high)
                unit_by_insertion[insertion] = meta.unit

            corr_s1_r, corr_s1_r2, corr_s1_n = _pearson_r2(series_by_insertion["S1"], series_by_insertion["S3"])
            corr_s2_r, corr_s2_r2, corr_s2_n = _pearson_r2(series_by_insertion["S2"], series_by_insertion["S3"])
            correlation_by_pair = {
                "S3-S1": {"r": corr_s1_r, "r2": corr_s1_r2, "n": corr_s1_n},
                "S3-S2": {"r": corr_s2_r, "r2": corr_s2_r2, "n": corr_s2_n},
            }
            transfer_confidence_by_pair = {
                "S3-S1": _normalized_transfer_metrics(series_by_insertion["S1"], series_by_insertion["S3"], limits_by_insertion["S3"][0], limits_by_insertion["S3"][1]),
                "S3-S2": _normalized_transfer_metrics(series_by_insertion["S2"], series_by_insertion["S3"], limits_by_insertion["S3"][0], limits_by_insertion["S3"][1]),
            }
            limit_mode = _limit_mode(limits_by_insertion["S3"][0], limits_by_insertion["S3"][1])
            if limit_mode == "Go-NoGo":
                analysis_path = "Go-NoGo Screening"
            elif limit_mode in {"Low-Only", "High-Only"}:
                analysis_path = "One-Sided Screening"
            else:
                analysis_path = "Parametric Guardband"

            s3_fail_mask = _fail_mask_for_values(series_by_insertion["S3"], limits_by_insertion["S3"][0], limits_by_insertion["S3"][1])

            current_screen_fail_masks = {
                "S1": _fail_mask_for_values(series_by_insertion["S1"], limits_by_insertion["S1"][0], limits_by_insertion["S1"][1]),
                "S2": _fail_mask_for_values(series_by_insertion["S2"], limits_by_insertion["S2"][0], limits_by_insertion["S2"][1]),
            }
            current_screen_recaptured_either_mask = current_screen_fail_masks["S1"] | current_screen_fail_masks["S2"]
            current_screen_metrics = {
                "recaptured_s1": int(np.count_nonzero(s3_fail_mask & current_screen_fail_masks["S1"])),
                "recaptured_s2": int(np.count_nonzero(s3_fail_mask & current_screen_fail_masks["S2"])),
                "recaptured_either": int(np.count_nonzero(s3_fail_mask & current_screen_recaptured_either_mask)),
                "residual_either": int(np.count_nonzero(s3_fail_mask & ~current_screen_recaptured_either_mask)),
                "residual_ppm_either": _ppm(int(np.count_nonzero(s3_fail_mask & ~current_screen_recaptured_either_mask)), screened_chip_count),
            }

            guardband_by_insertion: dict[str, dict[str, Any]] = {}
            guardband_fail_masks: dict[str, np.ndarray] = {}
            for insertion in ("S1", "S2"):
                guardband, median_delta, sigma_delta, paired_n = _guardband_shift(series_by_insertion[insertion], series_by_insertion["S3"])
                new_low, new_high, valid = _tighten_limits(limits_by_insertion[insertion][0], limits_by_insertion[insertion][1], guardband)
                guardband_by_insertion[insertion] = {
                    "guardband": guardband,
                    "median_delta": median_delta,
                    "sigma_delta": sigma_delta,
                    "paired_n": paired_n,
                    "new_low": new_low,
                    "new_high": new_high,
                    "valid": valid,
                }
                guardband_fail_masks[insertion] = _fail_mask_for_values(series_by_insertion[insertion], new_low, new_high) if valid else np.zeros(series_by_insertion[insertion].shape[0], dtype=bool)
            recaptured_s1 = int(np.count_nonzero(s3_fail_mask & guardband_fail_masks["S1"]))
            recaptured_s2 = int(np.count_nonzero(s3_fail_mask & guardband_fail_masks["S2"]))
            recaptured_either_mask = guardband_fail_masks["S1"] | guardband_fail_masks["S2"]
            recaptured_either = int(np.count_nonzero(s3_fail_mask & recaptured_either_mask))
            residual_either = int(np.count_nonzero(s3_fail_mask & ~recaptured_either_mask))

            overkill_s1 = int(np.count_nonzero(~s3_fail_mask & guardband_fail_masks["S1"]))
            overkill_s2 = int(np.count_nonzero(~s3_fail_mask & guardband_fail_masks["S2"]))
            overkill_either = int(np.count_nonzero(~s3_fail_mask & recaptured_either_mask))

            preferred_guardband = _select_preferred_guardband(recaptured_s1, _ppm(overkill_s1, screened_chip_count), corr_s1_r2, recaptured_s2, _ppm(overkill_s2, screened_chip_count), corr_s2_r2)
            adaptive_guardband_by_insertion = {
                "S1": _blank_adaptive_guardband(limits_by_insertion["S1"][0], limits_by_insertion["S1"][1]),
                "S2": _blank_adaptive_guardband(limits_by_insertion["S2"][0], limits_by_insertion["S2"][1]),
            }
            one_sided_screen_by_insertion = {
                "S1": _one_sided_screen_tighten(series_by_insertion["S1"], limits_by_insertion["S1"][0], limits_by_insertion["S1"][1], s3_fail_mask, screened_chip_count),
                "S2": _one_sided_screen_tighten(series_by_insertion["S2"], limits_by_insertion["S2"][0], limits_by_insertion["S2"][1], s3_fail_mask, screened_chip_count),
            }
            adaptive_guardband_used = False
            if analysis_path == "Go-NoGo Screening":
                preferred_guardband = _select_preferred_guardband(current_screen_metrics["recaptured_s1"], _ppm(int(np.count_nonzero(~s3_fail_mask & current_screen_fail_masks["S1"])), screened_chip_count), corr_s1_r2, current_screen_metrics["recaptured_s2"], _ppm(int(np.count_nonzero(~s3_fail_mask & current_screen_fail_masks["S2"])), screened_chip_count), corr_s2_r2)
                recaptured_s1 = current_screen_metrics["recaptured_s1"]
                recaptured_s2 = current_screen_metrics["recaptured_s2"]
                recaptured_either = current_screen_metrics["recaptured_either"]
                residual_either = current_screen_metrics["residual_either"]
                overkill_s1 = int(np.count_nonzero(~s3_fail_mask & current_screen_fail_masks["S1"]))
                overkill_s2 = int(np.count_nonzero(~s3_fail_mask & current_screen_fail_masks["S2"]))
                overkill_either = int(np.count_nonzero(~s3_fail_mask & current_screen_recaptured_either_mask))
            elif analysis_path == "One-Sided Screening":
                one_sided_candidates: list[tuple[str, float, float, float]] = []
                for insertion, r2_value in (("S1", corr_s1_r2), ("S2", corr_s2_r2)):
                    one_sided = one_sided_screen_by_insertion[insertion]
                    if one_sided["capture_all"]:
                        tighten_amount = one_sided["tighten_amount"] if math.isfinite(one_sided["tighten_amount"]) else float("inf")
                        overkill_ppm = one_sided["overkill_ppm"] if math.isfinite(one_sided["overkill_ppm"]) else float("inf")
                        one_sided_candidates.append((insertion, overkill_ppm, tighten_amount, -(r2_value if math.isfinite(r2_value) else -1.0)))
                if one_sided_candidates:
                    one_sided_candidates.sort(key=lambda item: (item[1], item[2], item[3], item[0]))
                    preferred_guardband = one_sided_candidates[0][0]
                recaptured_s1 = current_screen_metrics["recaptured_s1"]
                recaptured_s2 = current_screen_metrics["recaptured_s2"]
                recaptured_either = current_screen_metrics["recaptured_either"]
                residual_either = current_screen_metrics["residual_either"]
                overkill_s1 = int(np.count_nonzero(~s3_fail_mask & current_screen_fail_masks["S1"]))
                overkill_s2 = int(np.count_nonzero(~s3_fail_mask & current_screen_fail_masks["S2"]))
                overkill_either = int(np.count_nonzero(~s3_fail_mask & current_screen_recaptured_either_mask))
            elif np.count_nonzero(s3_fail_mask) > 0 and residual_either > 0:
                adaptive_guardband_by_insertion = {
                    "S1": _find_sigma_capture(series_by_insertion["S1"], series_by_insertion["S3"], limits_by_insertion["S1"][0], limits_by_insertion["S1"][1], s3_fail_mask, screened_chip_count),
                    "S2": _find_sigma_capture(series_by_insertion["S2"], series_by_insertion["S3"], limits_by_insertion["S2"][0], limits_by_insertion["S2"][1], s3_fail_mask, screened_chip_count),
                }

                capturing_candidates: list[tuple[str, float, float, float]] = []
                for insertion, r2_value in (("S1", corr_s1_r2), ("S2", corr_s2_r2)):
                    adaptive = adaptive_guardband_by_insertion[insertion]
                    if adaptive["capture_all"]:
                        sigma_value = adaptive["sigma_multiplier"] if math.isfinite(adaptive["sigma_multiplier"]) else float("inf")
                        overkill_ppm = adaptive["overkill_ppm"] if math.isfinite(adaptive["overkill_ppm"]) else float("inf")
                        capturing_candidates.append((insertion, overkill_ppm, sigma_value, -(r2_value if math.isfinite(r2_value) else -1.0)))
                if capturing_candidates:
                    capturing_candidates.sort(key=lambda item: (item[1], item[2], item[3], item[0]))
                    preferred_guardband = capturing_candidates[0][0]
                    adaptive_guardband_used = True

            escape_metrics = {
                "unique_escape_chips": int(np.count_nonzero(s3_fail_mask)),
                "unique_escape_ppm": _ppm(int(np.count_nonzero(s3_fail_mask)), screened_chip_count),
                "recaptured_s1": recaptured_s1,
                "recaptured_s2": recaptured_s2,
                "recaptured_either": recaptured_either,
                "residual_escape_either": residual_either,
                "residual_escape_ppm_either": _ppm(residual_either, screened_chip_count),
                "overkill_s1": overkill_s1,
                "overkill_s2": overkill_s2,
                "overkill_either": overkill_either,
                "overkill_s1_ppm": _ppm(overkill_s1, screened_chip_count),
                "overkill_s2_ppm": _ppm(overkill_s2, screened_chip_count),
                "overkill_either_ppm": _ppm(overkill_either, screened_chip_count),
                "preferred_guardband": preferred_guardband,
                "adaptive_guardband_used": adaptive_guardband_used,
            }

            plot_path = None
            if module_dir is not None:
                plot_path = _plot_test(
                    test=test,
                    module_dir=module_dir,
                    values_by_insertion=values_by_insertion,
                    stats_by_insertion=stats_by_insertion,
                    limits_by_insertion=limits_by_insertion,
                )
            report_rows.append(
                _collect_report_row(
                    test=test,
                    functional_block=functional_block,
                    plot_path=plot_path,
                    screened_chip_count=screened_chip_count,
                    analysis_path=analysis_path,
                    limit_mode=limit_mode,
                    stats_by_insertion=stats_by_insertion,
                    limits_by_insertion=limits_by_insertion,
                    unit_by_insertion=unit_by_insertion,
                    correlation_by_pair=correlation_by_pair,
                    transfer_confidence_by_pair=transfer_confidence_by_pair,
                    guardband_by_insertion=guardband_by_insertion,
                    adaptive_guardband_by_insertion=adaptive_guardband_by_insertion,
                    current_screen_metrics=current_screen_metrics,
                    one_sided_screen_by_insertion=one_sided_screen_by_insertion,
                    escape_metrics=escape_metrics,
                )
            )

            escape_indices = np.flatnonzero(s3_fail_mask)
            for idx in escape_indices:
                row_data = module_screened.iloc[idx]
                s1_value = series_by_insertion["S1"][idx]
                s2_value = series_by_insertion["S2"][idx]
                s3_value = series_by_insertion["S3"][idx]
                escape_detail_rows.append(
                    {
                        "Module": test.module,
                        "Test Number": int(test.test_number),
                        "Test Name": test.test_name,
                        "Chip Key": row_data["chip_key"],
                        "Wafer": row_data["WAFER"],
                        "X": row_data["X"],
                        "Y": row_data["Y"],
                        "S3 First Fail Test": row_data["FIRST_FAIL_TEST_S3"],
                        "S3 First Fail Name": row_data["S3_FIRST_FAIL_NAME"],
                        "S3 First Fail Module": row_data["S3_FIRST_FAIL_MODULE"],
                        "S3 First Fail Is Selected Module": "Y" if bool(row_data["S3_FIRST_FAIL_IN_SCOPE"]) else "N",
                        "Value S1": s1_value,
                        "Value S2": s2_value,
                        "Value S3": s3_value,
                        "LTL S1": limits_by_insertion["S1"][0],
                        "UTL S1": limits_by_insertion["S1"][1],
                        "LTL S2": limits_by_insertion["S2"][0],
                        "UTL S2": limits_by_insertion["S2"][1],
                        "LTL S3": limits_by_insertion["S3"][0],
                        "UTL S3": limits_by_insertion["S3"][1],
                        "Closest Margin S1": _closest_limit_margin(s1_value, limits_by_insertion["S1"][0], limits_by_insertion["S1"][1]),
                        "Closest Margin S2": _closest_limit_margin(s2_value, limits_by_insertion["S2"][0], limits_by_insertion["S2"][1]),
                        "Closest Margin S3": _closest_limit_margin(s3_value, limits_by_insertion["S3"][0], limits_by_insertion["S3"][1]),
                        "Guardband S1": guardband_by_insertion["S1"]["guardband"],
                        "Guardband S2": guardband_by_insertion["S2"]["guardband"],
                        "Recaptured by S1 Guardband": "Y" if guardband_fail_masks["S1"][idx] else "N",
                        "Recaptured by S2 Guardband": "Y" if guardband_fail_masks["S2"][idx] else "N",
                        "Recaptured by Either Guardband": "Y" if recaptured_either_mask[idx] else "N",
                    }
                )

            processed_tests += 1
            if processed_tests == 1 or processed_tests % 25 == 0 or processed_tests == total_tests_to_process:
                elapsed = perf_counter() - run_started
                pct = (100.0 * processed_tests / total_tests_to_process) if total_tests_to_process else 100.0
                eta_seconds = (elapsed / processed_tests) * (total_tests_to_process - processed_tests) if processed_tests else 0.0
                _log(
                    f"progress: {pct:6.2f}% ({processed_tests}/{total_tests_to_process} tests), elapsed={_format_duration(elapsed)}, eta={_format_duration(eta_seconds)}"
                )

        del frames
        del metas
        _log(f"[{module}] completed in {_format_duration(perf_counter() - module_started)}")

    report_df = pd.DataFrame(report_rows).sort_values(["Module", "Test Number"]).reset_index(drop=True)
    escape_detail_df = pd.DataFrame(escape_detail_rows).sort_values(["Module", "Test Number", "Wafer", "X", "Y"]).reset_index(drop=True) if escape_detail_rows else pd.DataFrame()
    _log("building workbook sheets")
    alignment_df = _build_alignment_summary_sheet(headers, raw_chip_frames, master_chip_df, selected_modules)
    module_pareto_df, test_pareto_df, ambient_escape_chips_df = _build_escape_pareto_sheets(master_chip_df, screened_chip_count)
    decision_matrix_df = _build_module_decision_matrix(report_df, escape_detail_df, screened_chip_count) if not report_df.empty else pd.DataFrame()
    skipped_df = pd.DataFrame(skipped_rows)
    metadata_df = _report_metadata(headers, common_tests, skipped_rows, master_chip_df, selected_modules, skip_plots, monitoring_patterns)
    analysis_path_metadata_df = _analysis_path_metadata(report_df)
    if not analysis_path_metadata_df.empty:
        metadata_df = pd.concat([metadata_df, analysis_path_metadata_df], ignore_index=True)
    recommendations_df = _recommended_analyses_sheet()

    _log("writing Excel workbook")
    with pd.ExcelWriter(REPORT_PATH, engine="openpyxl") as writer:
        alignment_df.to_excel(writer, sheet_name="Alignment_Summary", index=False)
        module_pareto_df.to_excel(writer, sheet_name="S3_Module_Pareto", index=False)
        test_pareto_df.to_excel(writer, sheet_name="S3_Test_Pareto", index=False)
        ambient_escape_chips_df.to_excel(writer, sheet_name="Ambient_Escape_Chips", index=False)
        report_df.to_excel(writer, sheet_name="Module_Test_Report", index=False)
        escape_detail_df.to_excel(writer, sheet_name="Module_Escape_Details", index=False)
        decision_matrix_df.to_excel(writer, sheet_name="Decision_Matrix", index=False)
        skipped_df.to_excel(writer, sheet_name="Skipped_Tests", index=False)
        recommendations_df.to_excel(writer, sheet_name="Recommended_Analyses", index=False)
        metadata_df.to_excel(writer, sheet_name="Report_Metadata", index=False)

    _log("styling workbook")
    _style_workbook(REPORT_PATH)
    _log(f"analysis completed in {_format_duration(perf_counter() - run_started)}")

    print(f"Workbook: {REPORT_PATH}")
    print(f"Plots root: {OUTPUT_DIR}")
    print(f"Selected modules: {', '.join(selected_modules)}")
    print(f"Comparable selected-module tests discovered: {len(common_tests)}")
    print(f"Comparable selected-module tests processed: {sum(len(tests) for tests in tests_by_module.values())}")
    print(f"Matched chips across S1/S2/S3: {len(master_chip_df)}")
    print(f"Screened chips passing S1 and S2: {screened_chip_count}")
    print(f"Ambient lot fails after S1/S2 screen: {int((master_chip_df['IS_SCREENED'] & master_chip_df['S3_LOT_FAIL']).sum())}")
    print(f"Skipped selected-module tests: {len(skipped_rows)}")
    print(f"Monitoring test-name patterns excluded: {', '.join(monitoring_patterns) if monitoring_patterns else 'None'}")
    print(f"Plots skipped: {'Y' if skip_plots else 'N'}")


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Analyze CTRX8188 production data for S3 ambient removal.")
    parser.add_argument(
        "--modules",
        nargs="*",
        default=list(RUN_MODULES),
        help="Modules of interest, e.g. DPLL TXPA TXLO or RXE2 REFX. Defaults to the TX module set for convenience.",
    )
    parser.add_argument(
        "--max-tests-per-module",
        type=int,
        default=RUN_MAX_TESTS_PER_MODULE,
        help="Optional limit used for smoke tests before a full run.",
    )
    parser.add_argument(
        "--skip-plots",
        action=argparse.BooleanOptionalAction,
        default=RUN_SKIP_PLOTS,
        help="Disable plot generation and only build the Excel report.",
    )
    parser.add_argument(
        "--monitoring-patterns",
        nargs="*",
        default=list(RUN_MONITORING_PATTERNS),
        help="Exclude tests whose names contain any of these substrings. Defaults include the TX monitoring patterns.",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = _parse_args()
    module_filter = _normalize_module_list(args.modules)
    run(selected_modules=module_filter, max_tests_per_module=args.max_tests_per_module, skip_plots=args.skip_plots, monitoring_patterns=tuple(args.monitoring_patterns))