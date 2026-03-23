from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[3]
WORK_DIR = REPO_ROOT / "Tasks_Automation_Code" / "IFX_Scripts" / "Correlation_Factors_FE_VNOM"
OUTPUT_DIR = WORK_DIR / "generated_outputs"

SUPPLY_VMIN = 0.95
SUPPLY_VNOM = 1.00
SUPPLY_VMAX = 1.05
INTERPOLATION_ALPHA = (SUPPLY_VNOM - SUPPLY_VMIN) / (SUPPLY_VMAX - SUPPLY_VMIN)

FACTOR_SHEET = "Correlation_Factors"
VNOM_COLUMN = "Corr_Factors_VNOM"
VMIN_COLUMN = "Corr_Factors_VMIN"
VMAX_COLUMN = "Corr_Factors_VMAX"

PLOT_EQUATION = (
    "Interpolation model:\n"
    "CF_VNOM = CF_VMIN + ((1.00 - 0.95) / (1.05 - 0.95)) * (CF_VMAX - CF_VMIN)\n"
    "CF_VNOM = 0.5 * (CF_VMIN + CF_VMAX)"
)

TEMPERATURE_LABELS = {
    -40: "Cold",
    25: "Ambient",
    135: "Hot",
}

TXLO_FREQUENCY_ORDER = [81, 77, 76]
TXLO_IDAC_ORDER = [10, 14, 18, 22, 27, 34, 42, 51, 63, 78, 96, 112]
TXPA_FREQUENCY_ORDER = [76, 77, 81]
TXPA_LUT_ORDER = [12, 22, 36, 48, 69, 86, 114, 130, 148, 170, 182, 195, 244]
TXPA_CHANNEL_ORDER = ["TX1", "TX2", "TX3", "TX4", "TX5", "TX6", "TX7", "TX8"]
TEXT_COLUMN_TO_CORNER = [(VMIN_COLUMN, "VMIN"), (VNOM_COLUMN, "VNOM"), (VMAX_COLUMN, "VMAX")]


@dataclass(frozen=True)
class WorkbookConfig:
    label: str
    source_path: Path
    group_columns: list[str]


WORKBOOKS = [
    WorkbookConfig(
        label="TXLO",
        source_path=WORK_DIR / "CV_ATE_Correlation_TXLO_Power_FE.xlsx",
        group_columns=["DataSheet", "Test Number", "Test Name", "Frequency_GHz", "LO IDAC", "N"],
    ),
    WorkbookConfig(
        label="TXPA",
        source_path=WORK_DIR / "CV_ATE_Correlation_TXPA_Power_FE.xlsx",
        group_columns=["DataSheet", "LUT value", "Test Number", "Test Name", "Frequency_GHz", "PA Channel"],
    ),
]


def _validate_factor_sheet(df: pd.DataFrame, workbook_name: str) -> None:
    missing = [column for column in [VMIN_COLUMN, VMAX_COLUMN] if column not in df.columns]
    if missing:
        raise ValueError(f"{workbook_name} is missing required columns: {missing}")
    if "Temperature" not in df.columns:
        raise ValueError(f"{workbook_name} is missing required column: Temperature")


def _add_vnom_column(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out[VMIN_COLUMN] = pd.to_numeric(out[VMIN_COLUMN], errors="coerce")
    out[VMAX_COLUMN] = pd.to_numeric(out[VMAX_COLUMN], errors="coerce")
    out[VNOM_COLUMN] = out[VMIN_COLUMN] + INTERPOLATION_ALPHA * (out[VMAX_COLUMN] - out[VMIN_COLUMN])
    return out


def _load_workbook_frames(path: Path) -> dict[str, pd.DataFrame]:
    xls = pd.ExcelFile(path)
    return {sheet_name: pd.read_excel(path, sheet_name=sheet_name) for sheet_name in xls.sheet_names}


def _copy_workbook_with_vnom(source_path: Path, factor_df: pd.DataFrame, output_path: Path) -> None:
    workbook = load_workbook(source_path)
    sheet = workbook[FACTOR_SHEET]

    header_row = 1
    header_map = {cell.value: cell.column for cell in sheet[header_row] if cell.value is not None}
    if VNOM_COLUMN in header_map:
        vnom_column_index = header_map[VNOM_COLUMN]
    else:
        vnom_column_index = sheet.max_column + 1
        sheet.cell(row=header_row, column=vnom_column_index, value=VNOM_COLUMN)

    for row_index, value in enumerate(factor_df[VNOM_COLUMN].tolist(), start=2):
        sheet.cell(row=row_index, column=vnom_column_index, value=None if pd.isna(value) else float(value))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _format_group_title(config: WorkbookConfig, group_key: tuple[object, ...]) -> str:
    parts = [f"{column}={value}" for column, value in zip(config.group_columns, group_key)]
    return f"{config.label} Correlation Factors | " + " | ".join(parts)


def _generate_pdf(config: WorkbookConfig, factor_df: pd.DataFrame, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    plot_df = factor_df.copy()
    plot_df["Temperature"] = pd.to_numeric(plot_df["Temperature"], errors="coerce")
    plot_df = plot_df.sort_values(config.group_columns + ["Temperature"])

    with PdfPages(output_path) as pdf:
        for group_key, group_df in plot_df.groupby(config.group_columns, dropna=False, sort=False):
            group_df = group_df.sort_values("Temperature")

            fig, ax = plt.subplots(figsize=(11.0, 6.5))
            ax.plot(group_df["Temperature"], group_df[VMIN_COLUMN], marker="o", linewidth=1.8, label="VMIN (95%)")
            ax.plot(group_df["Temperature"], group_df[VNOM_COLUMN], marker="s", linewidth=1.8, label="VNOM (100%)")
            ax.plot(group_df["Temperature"], group_df[VMAX_COLUMN], marker="^", linewidth=1.8, label="VMAX (105%)")

            ax.set_title(_format_group_title(config, group_key if isinstance(group_key, tuple) else (group_key,)))
            ax.set_xlabel("Temperature (C)")
            ax.set_ylabel("Correlation Factor")
            ax.grid(True, alpha=0.25)
            ax.legend(loc="best")
            ax.text(
                0.02,
                0.98,
                PLOT_EQUATION,
                transform=ax.transAxes,
                ha="left",
                va="top",
                fontsize=9,
                bbox=dict(boxstyle="round,pad=0.25", facecolor="white", alpha=0.85, edgecolor="none"),
            )

            fig.tight_layout()
            pdf.savefig(fig)
            plt.close(fig)


def _validate_interpolation(factor_df: pd.DataFrame, workbook_name: str) -> None:
    non_null = factor_df[[VMIN_COLUMN, VMAX_COLUMN, VNOM_COLUMN]].dropna()
    if non_null.empty:
        raise ValueError(f"{workbook_name} has no rows with usable VMIN/VMAX values")

    midpoint = 0.5 * (non_null[VMIN_COLUMN] + non_null[VMAX_COLUMN])
    max_abs_error = (non_null[VNOM_COLUMN] - midpoint).abs().max()
    if pd.isna(max_abs_error) or max_abs_error > 1e-12:
        raise ValueError(f"{workbook_name} interpolation validation failed; max abs error={max_abs_error}")


def _write_text_file(path: Path, values: list[float], decimals: int) -> None:
    formatted = [f"{float(value):.{decimals}f}" for value in values]
    path.write_text("\n".join(formatted) + "\n", encoding="ascii")


def _compare_existing_text(path: Path, values: list[float], decimals: int) -> None:
    expected = [f"{float(value):.{decimals}f}" for value in values]
    actual = [line.strip() for line in path.read_text(encoding="ascii").splitlines() if line.strip()]
    if actual != expected:
        raise ValueError(f"Existing text file does not match workbook-derived order/format: {path}")


def _export_txlo_texts(factor_df: pd.DataFrame) -> list[Path]:
    output_paths: list[Path] = []
    for temp, temp_label in TEMPERATURE_LABELS.items():
        subset = factor_df.loc[factor_df["Temperature"].eq(temp)].copy()
        subset["Frequency_GHz"] = pd.to_numeric(subset["Frequency_GHz"], errors="coerce")
        subset["LO IDAC"] = pd.to_numeric(subset["LO IDAC"], errors="coerce")
        subset["_freq_order"] = subset["Frequency_GHz"].map({value: idx for idx, value in enumerate(TXLO_FREQUENCY_ORDER)})
        subset["_idac_order"] = subset["LO IDAC"].map({value: idx for idx, value in enumerate(TXLO_IDAC_ORDER)})
        subset = subset.sort_values(["_freq_order", "_idac_order"])

        if len(subset) != len(TXLO_FREQUENCY_ORDER) * len(TXLO_IDAC_ORDER):
            raise ValueError(f"Unexpected TXLO row count for temperature {temp}: {len(subset)}")

        vmin_path = WORK_DIR / f"Corr_Factors_FE_TXLO_VMIN_81_77_76_GHz_{temp_label}.txt"
        _compare_existing_text(vmin_path, subset[VMIN_COLUMN].tolist(), decimals=3)

        for column_name, corner in TEXT_COLUMN_TO_CORNER:
            out_path = WORK_DIR / f"Corr_Factors_FE_TXLO_{corner}_81_77_76_GHz_{temp_label}.txt"
            _write_text_file(out_path, subset[column_name].tolist(), decimals=3)
            output_paths.append(out_path)

    return output_paths


def _export_txpa_texts(factor_df: pd.DataFrame) -> list[Path]:
    output_paths: list[Path] = []
    lut_order = {value: idx for idx, value in enumerate(TXPA_LUT_ORDER)}
    channel_order = {value: idx for idx, value in enumerate(TXPA_CHANNEL_ORDER)}

    for temp, temp_label in TEMPERATURE_LABELS.items():
        for freq in TXPA_FREQUENCY_ORDER:
            subset = factor_df.loc[factor_df["Temperature"].eq(temp) & factor_df["Frequency_GHz"].eq(freq)].copy()
            subset["LUT value"] = pd.to_numeric(subset["LUT value"], errors="coerce")
            subset["PA Channel"] = subset["PA Channel"].astype(str).str.strip().str.upper()

            base = subset.loc[subset["LUT value"].isin(TXPA_LUT_ORDER)].copy()
            base["_lut_order"] = base["LUT value"].map(lut_order)
            base = base.sort_values(["_lut_order"])

            tx255 = subset.loc[subset["LUT value"].eq(255)].copy()
            tx255["_channel_order"] = tx255["PA Channel"].map(channel_order)
            tx255 = tx255.sort_values(["_channel_order"])

            ordered = pd.concat([base, tx255], ignore_index=True)
            if len(ordered) != len(TXPA_LUT_ORDER) + len(TXPA_CHANNEL_ORDER):
                raise ValueError(f"Unexpected TXPA row count for temperature {temp}, frequency {freq}: {len(ordered)}")

            vmin_path = WORK_DIR / f"Corr_Factors_FE_TXPA_VMIN_{freq}GHz_{temp_label}.txt"
            _compare_existing_text(vmin_path, ordered[VMIN_COLUMN].tolist(), decimals=9)

            for column_name, corner in TEXT_COLUMN_TO_CORNER:
                out_path = WORK_DIR / f"Corr_Factors_FE_TXPA_{corner}_{freq}GHz_{temp_label}.txt"
                _write_text_file(out_path, ordered[column_name].tolist(), decimals=9)
                output_paths.append(out_path)

    return output_paths


def _export_text_files(config: WorkbookConfig, factor_df: pd.DataFrame) -> list[Path]:
    if config.label == "TXLO":
        return _export_txlo_texts(factor_df)
    if config.label == "TXPA":
        return _export_txpa_texts(factor_df)
    raise ValueError(f"Unsupported workbook label for text export: {config.label}")


def process_workbook(config: WorkbookConfig) -> tuple[Path, Path]:
    if not config.source_path.exists():
        raise FileNotFoundError(f"Missing source workbook: {config.source_path}")

    frames = _load_workbook_frames(config.source_path)
    if FACTOR_SHEET not in frames:
        raise ValueError(f"{config.source_path.name} is missing sheet: {FACTOR_SHEET}")

    factor_df = frames[FACTOR_SHEET]
    _validate_factor_sheet(factor_df, config.source_path.name)
    factor_df = _add_vnom_column(factor_df)
    _validate_interpolation(factor_df, config.source_path.name)

    workbook_output = OUTPUT_DIR / f"{config.source_path.stem}_VNOM.xlsx"
    pdf_output = OUTPUT_DIR / f"{config.label}_Correlation_Factors_All_Corners.pdf"

    _copy_workbook_with_vnom(config.source_path, factor_df, workbook_output)
    _generate_pdf(config, factor_df, pdf_output)
    text_outputs = _export_text_files(config, factor_df)
    for text_output in text_outputs:
        print(f"Wrote text: {text_output}")
    return workbook_output, pdf_output


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print(f"Interpolation alpha = {INTERPOLATION_ALPHA:.3f}")
    for config in WORKBOOKS:
        workbook_output, pdf_output = process_workbook(config)
        print(f"Wrote workbook: {workbook_output}")
        print(f"Wrote PDF: {pdf_output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())