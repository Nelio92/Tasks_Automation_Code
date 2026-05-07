"""Generate all-test correlation reports for flat IG-XL CSV datalogs.

The script expects the flat CSV format used by Test_Data_Reviewer:
row 1 contains metadata columns plus numeric test-number columns, row 2 maps
those test numbers to test names, optional limit/unit/stat rows follow, and
device measurement rows start at the first row whose first cell is numeric.
"""

from __future__ import annotations

import argparse
import csv
import math
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence


DEFAULT_ENCODING = "latin1"
DELIMITER = ";"

# User-configurable defaults. Command-line arguments can override these values.
SCRIPT_FOLDER = Path(__file__).resolve().parent
INPUT_FOLDER = SCRIPT_FOLDER
OUTPUT_FILE: Path | None = None

DEFAULT_THRESHOLD = 0.8
DEFAULT_MIN_PAIRED_VALUES = 3
DEFAULT_OUTPUT_NAME = "Test_Data_Correlation_Report.xlsx"
EXCEL_MAX_ROWS = 1_048_576

RESULT_HEADERS = [
    "Test Number 1",
    "Test Name 1",
    "Corr_Pearson",
    "Corr_Spearmann",
    "Rsquared_Pearson",
    "Rsquared_Spearmann",
    "Test Number 2",
    "Test Name 2",
    "Assessment",
]


@dataclass(frozen=True)
class FlatFileMeta:
    header: list[str]
    numeric_test_cols: list[str]
    data_start_line_index: int
    meta_rows: dict[str, dict[str, str]]


def _excel_col_letter(col_idx_1_based: int) -> str:
    if col_idx_1_based < 1:
        raise ValueError("Column index must be >= 1")
    n = col_idx_1_based
    letters: list[str] = []
    while n:
        n, rem = divmod(n - 1, 26)
        letters.append(chr(ord("A") + rem))
    return "".join(reversed(letters))


def _safe_sheet_name(name: str) -> str:
    safe = re.sub(r"[\\/*?:\[\]]", "_", name).strip() or "Sheet"
    return safe[:31]


def _unique_sheet_name(name: str, existing_names: Sequence[str]) -> str:
    base = _safe_sheet_name(name)
    existing_lower = {str(item).lower() for item in existing_names}
    if base.lower() not in existing_lower:
        return base

    counter = 2
    while True:
        suffix = f"_{counter}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}".strip() or f"Sheet{suffix}"
        if candidate.lower() not in existing_lower:
            return candidate
        counter += 1


def _autofit_openpyxl_columns(ws, *, min_width: int = 8, max_width: int = 70, padding: int = 2) -> None:
    if ws.max_column < 1 or ws.max_row < 1:
        return

    for col_idx, col_cells in enumerate(
        ws.iter_cols(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column),
        start=1,
    ):
        max_len = 0
        for cell in col_cells:
            if cell.value is None:
                continue
            text = str(cell.value)
            if "\n" in text:
                text = max(text.splitlines(), key=len)
            max_len = max(max_len, len(text))
        ws.column_dimensions[_excel_col_letter(col_idx)].width = max(min_width, min(max_width, max_len + padding))


def _shorten_test_name(test_name: str | None) -> str:
    raw = "" if test_name is None else str(test_name)
    cleaned = raw.strip().strip('"')
    if not cleaned:
        return ""
    primary, _, _ = cleaned.partition("<>")
    return primary.strip()


def _test_name_from_meta(meta: FlatFileMeta, test_col: str) -> str:
    return _shorten_test_name(meta.meta_rows.get("Test Name", {}).get(test_col))


def _to_excel_test_number(test_col: str) -> int | str:
    text = str(test_col).strip()
    if text.isdigit():
        return int(text)
    return text


def _clean_corr_value(value: object) -> float | None:
    try:
        corr = float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None
    if not math.isfinite(corr):
        return None
    if corr > 1.0 and corr <= 1.0 + 1e-12:
        return 1.0
    if corr < -1.0 and corr >= -1.0 - 1e-12:
        return -1.0
    return corr


def _r_squared(corr: float | None) -> float | None:
    if corr is None:
        return None
    return corr * corr


def _assessment(*values: float | None) -> str:
    finite_abs = [abs(value) for value in values if value is not None and math.isfinite(value)]
    if not finite_abs:
        return "no correlation"

    strength = max(finite_abs)
    if strength >= 0.8:
        return "strong correlation"
    if strength >= 0.5:
        return "moderate correlation"
    if strength >= 0.3:
        return "weak correlation"
    return "no correlation"


def _print_progress(label: str, current: int, total: int, *, end: str = "\n") -> None:
    if total <= 0:
        return

    clamped_current = min(max(current, 0), total)
    percent = clamped_current / total * 100.0
    print(f"\r  {label}: {percent:6.2f}% ({clamped_current}/{total})", end=end, flush=True)


def scan_flat_file_meta(
    file_path: Path,
    *,
    encoding: str = DEFAULT_ENCODING,
    delimiter: str = DELIMITER,
    needed_meta_rows: Iterable[str] = ("Test Name", "Low", "High", "Unit"),
    max_scan_lines: int = 200,
) -> FlatFileMeta:
    needed = {row_name.strip() for row_name in needed_meta_rows}
    meta_rows: dict[str, dict[str, str]] = {}

    with file_path.open("r", encoding=encoding, errors="replace", newline="") as stream:
        reader = csv.reader(stream, delimiter=delimiter)
        try:
            header = next(reader)
        except StopIteration as exc:
            raise ValueError(f"Empty file: {file_path}") from exc

        header = [item.strip() for item in header]
        numeric_test_cols = [item for item in header if item.strip().isdigit()]
        if not numeric_test_cols:
            raise ValueError("Could not find numeric test-number columns in the CSV header")

        data_start_line_index = 1
        for line_idx, row in enumerate(reader, start=1):
            if line_idx >= max_scan_lines:
                break
            if not row:
                continue
            key = (row[0] or "").strip().strip('"')
            if key.isdigit():
                data_start_line_index = line_idx
                break
            if key in needed:
                row_map: dict[str, str] = {}
                for col_name, cell in zip(header, row, strict=False):
                    if col_name in numeric_test_cols:
                        row_map[col_name] = (cell or "").strip()
                meta_rows[key] = row_map
        else:
            data_start_line_index = max_scan_lines

    return FlatFileMeta(
        header=header,
        numeric_test_cols=numeric_test_cols,
        data_start_line_index=data_start_line_index,
        meta_rows=meta_rows,
    )


def _read_measurement_data(file_path: Path, meta: FlatFileMeta, *, encoding: str):
    import numpy as np
    import pandas as pd

    df = pd.read_csv(
        file_path,
        sep=DELIMITER,
        encoding=encoding,
        low_memory=False,
        usecols=meta.numeric_test_cols,
        skiprows=range(1, meta.data_start_line_index),
        decimal=",",
        memory_map=True,
    )
    df.columns = [str(col).strip() for col in df.columns]

    object_cols = df.select_dtypes(include="object").columns
    for col in object_cols:
        df[col] = df[col].astype(str).str.strip().str.strip('"').str.replace(",", ".", regex=False)

    df = df.apply(pd.to_numeric, errors="coerce")
    return df.replace([np.inf, -np.inf], np.nan)


def _drop_non_correlatable_tests(df, *, min_paired_values: int):
    counts = df.count()
    unique_counts = df.nunique(dropna=True)
    valid_cols = [
        col
        for col in df.columns
        if int(counts.get(col, 0)) >= min_paired_values and int(unique_counts.get(col, 0)) >= 2
    ]
    return df.loc[:, valid_cols]


def _append_correlation_rows(
    ws,
    *,
    meta: FlatFileMeta,
    df,
    threshold: float,
    min_paired_values: int,
    max_rows_per_sheet: int,
) -> tuple[int, bool]:
    import numpy as np

    _print_progress("Current file progress", 0, 100)
    df = _drop_non_correlatable_tests(df, min_paired_values=min_paired_values)
    if df.shape[1] < 2:
        _print_progress("Current file progress", 100, 100)
        return 0, False

    print(f"  Correlatable tests: {df.shape[1]}")
    _print_progress("Current file progress", 10, 100)
    print("  Calculating Pearson matrix...")
    pearson_corr = df.corr(method="pearson", min_periods=min_paired_values)
    _print_progress("Current file progress", 35, 100)
    print("  Ranking values for Spearman matrix...")
    ranked_df = df.rank(axis=0, method="average", na_option="keep")
    _print_progress("Current file progress", 45, 100)
    print("  Calculating Spearman matrix...")
    spearman_corr = ranked_df.corr(method="pearson", min_periods=min_paired_values)
    _print_progress("Current file progress", 70, 100)
    print("  Filtering strong pairs...")

    pearson_values = pearson_corr.to_numpy(dtype=float)
    spearman_values = spearman_corr.to_numpy(dtype=float)
    strong_mask = (np.abs(pearson_values) >= threshold) | (np.abs(spearman_values) >= threshold)
    strong_mask &= np.triu(np.ones(strong_mask.shape, dtype=bool), k=1)
    row_indices, col_indices = np.where(strong_mask)
    if row_indices.size == 0:
        _print_progress("Current file progress", 100, 100)
        return 0, False

    pearson_pair_values = pearson_values[row_indices, col_indices]
    spearman_pair_values = spearman_values[row_indices, col_indices]
    strengths = np.maximum(
        np.nan_to_num(np.abs(pearson_pair_values), nan=-np.inf),
        np.nan_to_num(np.abs(spearman_pair_values), nan=-np.inf),
    )
    order = np.argsort(-strengths, kind="stable")
    _print_progress("Current file progress", 80, 100)

    columns = list(df.columns)
    output_rows = 0
    truncated = False
    rows_to_write = min(int(order.size), max_rows_per_sheet)
    last_progress_percent = -1
    _print_progress("Writing Excel rows", 0, rows_to_write, end="")
    for order_idx in order:
        if output_rows >= max_rows_per_sheet:
            truncated = True
            break

        left_idx = int(row_indices[order_idx])
        right_idx = int(col_indices[order_idx])
        left_col = columns[left_idx]
        right_col = columns[right_idx]
        pearson_value = _clean_corr_value(pearson_values[left_idx, right_idx])
        spearman_value = _clean_corr_value(spearman_values[left_idx, right_idx])

        ws.append(
            [
                _to_excel_test_number(left_col),
                _test_name_from_meta(meta, left_col),
                pearson_value,
                spearman_value,
                _r_squared(pearson_value),
                _r_squared(spearman_value),
                _to_excel_test_number(right_col),
                _test_name_from_meta(meta, right_col),
                _assessment(pearson_value, spearman_value),
            ]
        )
        output_rows += 1

        current_progress_percent = int(output_rows * 100 / rows_to_write)
        if current_progress_percent != last_progress_percent or output_rows == rows_to_write:
            _print_progress("Writing Excel rows", output_rows, rows_to_write, end="")
            last_progress_percent = current_progress_percent

    print()
    _print_progress("Current file progress", 100, 100)

    return output_rows, truncated


def _style_sheet(ws) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor="1F4E78")
    strong_fill = PatternFill("solid", fgColor="D9EAD3")
    moderate_fill = PatternFill("solid", fgColor="FFF2CC")
    weak_fill = PatternFill("solid", fgColor="FCE4D6")
    no_fill = PatternFill("solid", fgColor="E7E6E6")

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in ws.iter_rows(min_row=2, min_col=3, max_col=6):
        for cell in row:
            cell.number_format = "0.0000"

    assessment_col = RESULT_HEADERS.index("Assessment") + 1
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=assessment_col)
        if cell.value == "strong correlation":
            cell.fill = strong_fill
        elif cell.value == "moderate correlation":
            cell.fill = moderate_fill
        elif cell.value == "weak correlation":
            cell.fill = weak_fill
        elif cell.value == "no correlation":
            cell.fill = no_fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{_excel_col_letter(ws.max_column)}{ws.max_row}"
    _autofit_openpyxl_columns(ws)


def _collect_csv_paths(
    input_folder: Path,
    *,
    pattern: str,
    single_file: str | None,
    max_files: int | None,
) -> list[Path]:
    if single_file:
        candidate = Path(single_file)
        csv_paths = [candidate if candidate.is_absolute() else input_folder / candidate]
    else:
        csv_paths = sorted(path for path in input_folder.glob(pattern) if path.is_file() and path.suffix.lower() == ".csv")

    if max_files is not None:
        csv_paths = csv_paths[:max_files]
    return csv_paths


def generate_correlation_workbook(
    *,
    input_folder: Path,
    output_file: Path,
    threshold: float = DEFAULT_THRESHOLD,
    min_paired_values: int = DEFAULT_MIN_PAIRED_VALUES,
    encoding: str = DEFAULT_ENCODING,
    pattern: str = "*.csv",
    single_file: str | None = None,
    max_files: int | None = None,
    max_rows_per_sheet: int = EXCEL_MAX_ROWS - 1,
) -> Path:
    from openpyxl import Workbook

    if threshold < 0.0 or threshold > 1.0:
        raise ValueError("threshold must be between 0.0 and 1.0")
    if min_paired_values < 2:
        raise ValueError("min_paired_values must be at least 2")
    if max_rows_per_sheet < 1 or max_rows_per_sheet > EXCEL_MAX_ROWS - 1:
        raise ValueError(f"max_rows_per_sheet must be between 1 and {EXCEL_MAX_ROWS - 1}")

    csv_paths = _collect_csv_paths(input_folder, pattern=pattern, single_file=single_file, max_files=max_files)
    if not csv_paths:
        raise SystemExit(f"No CSV files found in: {input_folder}")

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    total_files = len(csv_paths)
    for file_idx, file_path in enumerate(csv_paths, start=1):
        print(f"[{file_idx}/{total_files}] Processing {file_path.name}")
        _print_progress("Overall file progress", file_idx - 1, total_files)
        ws = wb.create_sheet(_unique_sheet_name(file_path.stem, wb.sheetnames))
        ws.append(RESULT_HEADERS)

        try:
            meta = scan_flat_file_meta(file_path, encoding=encoding)
            df = _read_measurement_data(file_path, meta, encoding=encoding)
            out_rows, truncated = _append_correlation_rows(
                ws,
                meta=meta,
                df=df,
                threshold=threshold,
                min_paired_values=min_paired_values,
                max_rows_per_sheet=max_rows_per_sheet,
            )
        except Exception as exc:
            ws.append(["ERROR", str(exc), None, None, None, None, None, None, "no correlation"])
            out_rows = 1
            truncated = False
            print(f"  ERROR: {exc}")
        else:
            trunc_msg = " (truncated at configured row limit)" if truncated else ""
            print(f"  Recorded {out_rows} correlation pair(s){trunc_msg}")

        _style_sheet(ws)
        _print_progress("Overall file progress", file_idx, total_files)

    try:
        wb.save(output_file)
    except PermissionError:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        fallback = output_file.with_name(f"{output_file.stem}_{timestamp}{output_file.suffix}")
        wb.save(fallback)
        print(f"Could not overwrite open workbook: {output_file}")
        print(f"Saved instead: {fallback}")
        return fallback

    print(f"Saved: {output_file}")
    return output_file


def _parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create an Excel workbook of strong Pearson/Spearman correlations for flat CSV datalogs.",
    )
    parser.add_argument(
        "--input-folder",
        type=Path,
        default=INPUT_FOLDER,
        help=f"Folder containing flat CSV datalogs. Default: {INPUT_FOLDER}.",
    )
    parser.add_argument(
        "--output-file",
        type=Path,
        default=OUTPUT_FILE,
        help="Excel output path. Default: <input-folder>/Outputs/Test_Data_Correlation_Report.xlsx.",
    )
    parser.add_argument(
        "--threshold",
        type=float,
        default=DEFAULT_THRESHOLD,
        help="Absolute correlation threshold for recording rows. Default: 0.8.",
    )
    parser.add_argument(
        "--min-paired-values",
        type=int,
        default=DEFAULT_MIN_PAIRED_VALUES,
        help="Minimum paired numeric values needed to calculate a correlation. Default: 3.",
    )
    parser.add_argument("--encoding", default=DEFAULT_ENCODING, help="CSV text encoding. Default: latin1.")
    parser.add_argument("--pattern", default="*.csv", help="Input glob pattern within --input-folder. Default: *.csv.")
    parser.add_argument("--single-file", default=None, help="Analyze only one CSV file name/path.")
    parser.add_argument("--max-files", type=int, default=None, help="Analyze at most this many CSV files.")
    parser.add_argument(
        "--max-rows-per-sheet",
        type=int,
        default=EXCEL_MAX_ROWS - 1,
        help="Maximum data rows written per worksheet. Default: Excel row limit minus header.",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = _parse_args(argv)
    input_folder = args.input_folder.resolve()
    output_file = args.output_file
    if output_file is None:
        output_file = input_folder / "Outputs" / DEFAULT_OUTPUT_NAME
    else:
        output_file = output_file.resolve()

    generate_correlation_workbook(
        input_folder=input_folder,
        output_file=output_file,
        threshold=args.threshold,
        min_paired_values=args.min_paired_values,
        encoding=args.encoding,
        pattern=args.pattern,
        single_file=args.single_file,
        max_files=args.max_files,
        max_rows_per_sheet=args.max_rows_per_sheet,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
