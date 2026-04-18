from __future__ import annotations

import sys
import tempfile
import textwrap
import unittest
from pathlib import Path

from openpyxl import load_workbook


TEST_DATA_REVIEWER_DIR = Path(__file__).resolve().parents[1]
if str(TEST_DATA_REVIEWER_DIR) not in sys.path:
    sys.path.insert(0, str(TEST_DATA_REVIEWER_DIR))

import Test_Data_Reviewer as analysis


class ReviewDatasetWaterfallCleanupTests(unittest.TestCase):
    def test_collect_review_dataset_discovers_modules_and_applies_waterfall_cleanup(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_root = Path(tmp_dir)
            input_folder = tmp_root / "input"
            output_folder = tmp_root / "output"
            input_folder.mkdir()
            output_folder.mkdir()

            csv_path = input_folder / "waterfall_sample.csv"
            csv_path.write_text(
                textwrap.dedent(
                    """\
                    UNIT_ID;SITE_NUM;WAFER;X;Y;LOT;CHIP_ID;PF;FIRST_FAIL_TEST;1001;1999;2001
                    Test Name;;;;;;;;;MODA_GAIN_MAIN;TIME_MODA_FUNC_S980;MODB_GAIN_MAIN
                    Low;;;;;;;;;0.5;0.0;0.5
                    High;;;;;;;;;2.0;100.0;2.0
                    Unit;;;;;;;;;dB;ms;dB
                    Cpk;;;;;;;;;1.0;5.0;1.0
                    Yield;;;;;;;;;66.7;100.0;66.7
                    Mean;;;;;;;;;0.7;10.0;2.3
                    Stddev;;;;;;;;;0.6;1.0;1.9
                    1;0;W1;1;1;LOT_A;C1;F;MODA_GAIN_MAIN;0.0;10.0;5.0
                    2;0;W1;1;2;LOT_A;C2;P;;1.0;11.0;1.0
                    3;0;W1;1;3;LOT_A;C3;P;;1.1;12.0;1.1
                    """
                ),
                encoding=analysis.DEFAULT_ENCODING,
            )

            without_cleanup = analysis.collect_review_dataset(
                input_folder=input_folder,
                output_folder=output_folder,
                modules=None,
                outlier_mad_multiplier=6.0,
                yield_threshold=100.0,
                cpk_low=1.67,
                cpk_high=20.0,
                max_files=None,
                single_file=None,
                selected_csv_paths=[csv_path],
                waterfall_cleanup_enabled=False,
            )
            with_cleanup = analysis.collect_review_dataset(
                input_folder=input_folder,
                output_folder=output_folder,
                modules=None,
                outlier_mad_multiplier=6.0,
                yield_threshold=100.0,
                cpk_low=1.67,
                cpk_high=20.0,
                max_files=None,
                single_file=None,
                selected_csv_paths=[csv_path],
                waterfall_cleanup_enabled=True,
            )

            self.assertEqual(without_cleanup.available_modules, ("MODA", "MODB"))
            self.assertEqual(with_cleanup.available_modules, ("MODA", "MODB"))
            self.assertFalse(without_cleanup.waterfall_cleanup_enabled)
            self.assertTrue(with_cleanup.waterfall_cleanup_enabled)

            self.assertEqual([finding.test_name for finding in without_cleanup.findings], ["MODA_GAIN_MAIN", "MODB_GAIN_MAIN"])
            self.assertEqual([finding.test_name for finding in with_cleanup.findings], ["MODA_GAIN_MAIN"])

            moda_finding = with_cleanup.findings[0]
            self.assertEqual(moda_finding.module, "MODA")
            self.assertEqual(moda_finding.fail_chips, 1)
            self.assertAlmostEqual(moda_finding.yield_pct or 0.0, 66.6666666667, places=3)
            self.assertEqual(moda_finding.cleanup_removed_before_review, 0)

            cleanup_summaries = with_cleanup.cleanup_summaries_by_file[csv_path.name]
            self.assertEqual(len(cleanup_summaries), 2)
            self.assertEqual(cleanup_summaries[0].removed_chip_count, 1)
            self.assertEqual(cleanup_summaries[0].remaining_chip_count, 2)
            self.assertEqual(cleanup_summaries[1].removed_chip_count, 0)
            self.assertEqual(cleanup_summaries[1].remaining_chip_count, 2)

    def test_waterfall_cleanup_is_scoped_per_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_root = Path(tmp_dir)
            input_folder = tmp_root / "input"
            output_folder = tmp_root / "output"
            input_folder.mkdir()
            output_folder.mkdir()

            first_path = input_folder / "first.csv"
            first_path.write_text(
                textwrap.dedent(
                    """\
                    UNIT_ID;SITE_NUM;WAFER;X;Y;LOT;CHIP_ID;PF;FIRST_FAIL_TEST;1001;1999;2001
                    Test Name;;;;;;;;;MODA_GAIN_MAIN;TIME_MODA_FUNC_S980;MODB_GAIN_MAIN
                    Low;;;;;;;;;0.5;0.0;0.5
                    High;;;;;;;;;2.0;100.0;2.0
                    Unit;;;;;;;;;dB;ms;dB
                    Cpk;;;;;;;;;1.0;5.0;1.0
                    Yield;;;;;;;;;66.7;100.0;66.7
                    Mean;;;;;;;;;0.7;10.0;2.3
                    Stddev;;;;;;;;;0.6;1.0;1.9
                    1;0;W1;1;1;LOT_A;A1;F;MODA_GAIN_MAIN;0.0;10.0;5.0
                    2;0;W1;1;2;LOT_A;A2;P;;1.0;11.0;1.0
                    3;0;W1;1;3;LOT_A;A3;P;;1.1;12.0;1.1
                    """
                ),
                encoding=analysis.DEFAULT_ENCODING,
            )

            second_path = input_folder / "second.csv"
            second_path.write_text(
                textwrap.dedent(
                    """\
                    UNIT_ID;SITE_NUM;WAFER;X;Y;LOT;CHIP_ID;PF;FIRST_FAIL_TEST;1001;1999;2001
                    Test Name;;;;;;;;;MODA_GAIN_MAIN;TIME_MODA_FUNC_S980;MODB_GAIN_MAIN
                    Low;;;;;;;;;0.5;0.0;0.5
                    High;;;;;;;;;2.0;100.0;2.0
                    Unit;;;;;;;;;dB;ms;dB
                    Cpk;;;;;;;;;1.0;5.0;1.0
                    Yield;;;;;;;;;100.0;100.0;66.7
                    Mean;;;;;;;;;1.0;10.0;2.3
                    Stddev;;;;;;;;;0.1;1.0;1.9
                    1;0;W1;1;1;LOT_B;B1;P;;1.0;10.0;5.0
                    2;0;W1;1;2;LOT_B;B2;P;;1.0;11.0;1.0
                    3;0;W1;1;3;LOT_B;B3;P;;1.1;12.0;1.1
                    """
                ),
                encoding=analysis.DEFAULT_ENCODING,
            )

            dataset = analysis.collect_review_dataset(
                input_folder=input_folder,
                output_folder=output_folder,
                modules=None,
                outlier_mad_multiplier=6.0,
                yield_threshold=100.0,
                cpk_low=1.67,
                cpk_high=20.0,
                max_files=None,
                single_file=None,
                selected_csv_paths=[first_path, second_path],
                waterfall_cleanup_enabled=True,
            )

            findings = [(finding.file_name, finding.test_name) for finding in dataset.findings]
            self.assertEqual(
                findings,
                [
                    ("first.csv", "MODA_GAIN_MAIN"),
                    ("second.csv", "MODB_GAIN_MAIN"),
                ],
            )

    def test_waterfall_cleanup_uses_function_sequence_not_module_name(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_root = Path(tmp_dir)
            input_folder = tmp_root / "input"
            output_folder = tmp_root / "output"
            input_folder.mkdir()
            output_folder.mkdir()

            csv_path = input_folder / "function_sequence.csv"
            csv_path.write_text(
                textwrap.dedent(
                    """\
                    UNIT_ID;SITE_NUM;WAFER;X;Y;LOT;CHIP_ID;PF;FIRST_FAIL_TEST;1001;1999;2001;2999;3001
                    Test Name;;;;;;;;;MODA_GAIN_STAGE1;TIME_MODA_F1_S980;MODB_GAIN_STAGE1;TIME_MODB_F2_S980;MODA_GAIN_STAGE2
                    Low;;;;;;;;;0.5;0.0;0.5;0.0;0.5
                    High;;;;;;;;;2.0;100.0;2.0;100.0;2.0
                    Unit;;;;;;;;;dB;ms;dB;ms;dB
                    Cpk;;;;;;;;;1.0;5.0;1.0;5.0;1.0
                    Yield;;;;;;;;;66.7;100.0;66.7;100.0;100.0
                    Mean;;;;;;;;;0.7;10.0;2.3;10.0;2.3
                    Stddev;;;;;;;;;0.6;1.0;1.9;1.0;1.9
                    1;0;W1;1;1;LOT_A;C1;F;MODA_GAIN_STAGE1;0.0;10.0;5.0;10.0;5.0
                    2;0;W1;1;2;LOT_A;C2;F;MODB_GAIN_STAGE1;1.0;11.0;5.0;11.0;5.0
                    3;0;W1;1;3;LOT_A;C3;F;MODA_GAIN_STAGE2;1.1;12.0;1.1;12.0;5.0
                    """
                ),
                encoding=analysis.DEFAULT_ENCODING,
            )

            dataset = analysis.collect_review_dataset(
                input_folder=input_folder,
                output_folder=output_folder,
                modules=["MODA", "MODB"],
                outlier_mad_multiplier=6.0,
                yield_threshold=100.0,
                cpk_low=1.67,
                cpk_high=20.0,
                max_files=None,
                single_file=None,
                selected_csv_paths=[csv_path],
                waterfall_cleanup_enabled=True,
            )

            findings = [(finding.test_name, finding.fail_chips, round(finding.yield_pct or 0.0, 3)) for finding in dataset.findings]
            self.assertEqual(
                findings,
                [
                    ("MODA_GAIN_STAGE1", 1, 66.667),
                    ("MODB_GAIN_STAGE1", 1, 50.0),
                    ("MODA_GAIN_STAGE2", 1, 0.0),
                ],
            )

            self.assertEqual([finding.cleanup_removed_before_review for finding in dataset.findings], [0, 1, 2])
            self.assertEqual([finding.sample_count for finding in dataset.findings], [3, 2, 1])

            cleanup_summaries = dataset.cleanup_summaries_by_file[csv_path.name]
            self.assertEqual(
                [(item.function_index, item.removed_chip_count, item.remaining_chip_count) for item in cleanup_summaries],
                [(1, 1, 2), (2, 1, 1), (3, 1, 0)],
            )

    def test_export_workbook_includes_cleanup_summary_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_root = Path(tmp_dir)
            input_folder = tmp_root / "input"
            output_folder = tmp_root / "output"
            input_folder.mkdir()
            output_folder.mkdir()

            csv_path = input_folder / "cleanup_export.csv"
            csv_path.write_text(
                textwrap.dedent(
                    """\
                    UNIT_ID;SITE_NUM;WAFER;X;Y;LOT;CHIP_ID;PF;FIRST_FAIL_TEST;1001;1999;2001;2999;3001
                    Test Name;;;;;;;;;MODA_GAIN_STAGE1;TIME_MODA_F1_S980;MODB_GAIN_STAGE1;TIME_MODB_F2_S980;MODA_GAIN_STAGE2
                    Low;;;;;;;;;0.5;0.0;0.5;0.0;0.5
                    High;;;;;;;;;2.0;100.0;2.0;100.0;2.0
                    Unit;;;;;;;;;dB;ms;dB;ms;dB
                    Cpk;;;;;;;;;1.0;5.0;1.0;5.0;1.0
                    Yield;;;;;;;;;66.7;100.0;66.7;100.0;100.0
                    Mean;;;;;;;;;0.7;10.0;2.3;10.0;2.3
                    Stddev;;;;;;;;;0.6;1.0;1.9;1.0;1.9
                    1;0;W1;1;1;LOT_A;C1;F;MODA_GAIN_STAGE1;0.0;10.0;5.0;10.0;5.0
                    2;0;W1;1;2;LOT_A;C2;F;MODB_GAIN_STAGE1;1.0;11.0;5.0;11.0;5.0
                    3;0;W1;1;3;LOT_A;C3;F;MODA_GAIN_STAGE2;1.1;12.0;1.1;12.0;5.0
                    """
                ),
                encoding=analysis.DEFAULT_ENCODING,
            )

            dataset = analysis.collect_review_dataset(
                input_folder=input_folder,
                output_folder=output_folder,
                modules=["MODA", "MODB"],
                outlier_mad_multiplier=6.0,
                yield_threshold=100.0,
                cpk_low=1.67,
                cpk_high=20.0,
                max_files=None,
                single_file=None,
                selected_csv_paths=[csv_path],
                waterfall_cleanup_enabled=True,
            )

            workbook_path = analysis.export_review_dataset_workbook(
                dataset,
                destination_path=output_folder / "cleanup_report.xlsx",
            )

            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                self.assertIn("Cleanup Summary", workbook.sheetnames)
                cleanup_worksheet = workbook["Cleanup Summary"]
                self.assertEqual(cleanup_worksheet["A1"].value, "Waterfall Cleanup Summary")
                self.assertEqual(cleanup_worksheet["A3"].value, "Waterfall cleanup")
                self.assertEqual(cleanup_worksheet["B3"].value, "Enabled")
                self.assertEqual(cleanup_worksheet["A7"].value, "Total removed chips")
                self.assertEqual(cleanup_worksheet["B7"].value, 3)
                self.assertEqual(cleanup_worksheet["A10"].value, "Function-by-function cleanup")
                self.assertEqual(cleanup_worksheet["A11"].value, "File")
                self.assertEqual(cleanup_worksheet["E11"].value, "Removed Chips")
                self.assertEqual(cleanup_worksheet["A12"].value, csv_path.name)
                self.assertEqual(cleanup_worksheet["B12"].value, 1)
                self.assertEqual(cleanup_worksheet["C12"].value, "TIME_MODA_F1_S980")
                self.assertEqual(cleanup_worksheet["E12"].value, 1)
                self.assertEqual(cleanup_worksheet["F14"].value, 0)
                self.assertEqual(cleanup_worksheet["G12"].value, "Open")
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()